# -*- coding: utf-8 -*-
"""Build the Anger chapter (handwritten sheets 1-11) as a Word doc matching the
style of 'Introduction - final.docx' and the other chapters: Garamond, A5, 13pt
justified. References in plain parentheses, hyphens in verse ranges, no em dashes.
All handwritten English; references are cited inline (no block quotations)."""
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Garamond"
BODY_SZ = Pt(13)

doc = Document()

# ---- base style ----
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = BODY_SZ
rpr = normal.element.get_or_add_rPr()
rf = rpr.get_or_add_rFonts()
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

# ---- page = A5 ----
sec = doc.sections[0]
sec.page_width = Mm(148)
sec.page_height = Mm(210)
sec.top_margin = Mm(14)
sec.bottom_margin = Mm(15)
sec.left_margin = Mm(15)
sec.right_margin = Mm(13)

# ---- footer page number ----
def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p)

# ---- helpers ----
def _set_run(r, bold=False, italic=False, size=None):
    r.font.name = FONT
    rpr = r._r.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    r.font.bold = bold
    r.font.italic = italic
    if size: r.font.size = size

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(text.upper())
    _set_run(r, bold=True, size=Pt(22))
    return p

def section_heading(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(5)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(15))
    return p

def lead(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text); _set_run(r)
    return p

def bullet(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("•  "); _set_run(r)
    r = p.add_run(text); _set_run(r)
    return p

def bullet2(lab, rest):
    """Bullet whose leading label is bold, then normal text."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("•  "); _set_run(r)
    r = p.add_run(lab); _set_run(r, bold=True)
    r = p.add_run(rest); _set_run(r)
    return p

def subbullet(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(12)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("–  "); _set_run(r)
    r = p.add_run(text); _set_run(r)
    return p

def riddle(lines):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    for i, ln in enumerate(lines):
        r = p.add_run(ln); _set_run(r, italic=True)
        if i != len(lines) - 1:
            r.add_break()
    return p

# ================= CONTENT =================

title("Anger")

riddle([
    "Riddle! Riddle!",
    "I am a five-lettered word.",
    "I am a giant inside your chest.",
    "I make you shout and slam the door.",
    "The longer you make me stay and rest,",
    "the more I leave you bruised and sore.",
    "Who am I?",
])

# ---- What is anger? ----
section_heading("What is anger?")
lead("In the Bible, anger is a powerful emotion that is not always sinful. It "
     "functions as a natural human response to perceived wrongs or injustices, "
     "but it becomes dangerous and sinful when it is driven by pride, turns into "
     "bitterness, or leads to destructive actions.")

# ---- Types of anger ----
section_heading("Types of anger")
bullet2("Righteous anger: ", "a holy displeasure against sin, injustice, and "
        "evil. God's anger is a perfect, just reaction to evil.")
subbullet("Jesus displayed righteous anger when clearing the corrupt merchants "
          "out of the temple (John 2:13-18).")
bullet2("Sinful anger: ", "a human anger rooted in selfishness, pride, or a "
        "desire for revenge.")
subbullet("This includes sudden outbursts of rage, holding grudges, or malicious "
          "shouting and hostility.")

# ---- Biblical principles ----
section_heading("What are the biblical principles of anger?")
bullet2("Anger is manageable. ", "Scripture notes, “In your anger do not sin” "
        "(Ephesians 4:26), acknowledging that the feeling exists while commending "
        "self-control.")
bullet2("Deal with anger quickly. ", "Ephesians 4:26 also warns not to let the "
        "sun go down on your anger, meaning you should resolve issues promptly so "
        "that bitterness does not take root.")
bullet2("Be slow to anger. ", "James 1:19-20 advises believers to be quick to "
        "listen and slow to become angry, noting that human anger does not achieve "
        "God's righteous standards.")
bullet2("Leave vengeance to God. ", "Romans 12:19 instructs Christians not to "
        "take revenge into their own hands, but to trust in God's ultimate "
        "justice.")

# ---- What brings God's anger? ----
section_heading("What brings God's anger?")
lead("God's anger is a righteous and loving response to evil, injustice, and "
     "actions that destroy the people He cares about. According to biblical "
     "teaching, several specific attitudes and actions provoke divine "
     "displeasure. These include:")
bullet2("Pride. ", "Believing you do not need God, acting with arrogance, or "
        "looking down on other people.")
bullet2("Injustice and oppression. ", "Mistreating the poor, the weak, or the "
        "vulnerable, or shedding innocent blood.")
bullet2("Lies and deceit. ", "Spreading falsehoods, twisting the truth, or "
        "bearing false witness.")
bullet2("Wicked intentions. ", "Plotting evil in secret, or enjoying deliberate "
        "rebellion against what is good and acceptable to God and His people.")
bullet2("Sowing division. ", "Causing arguments, conflict, and deep division "
        "among people.")
lead("Wrong attitudes and actions towards God's chosen ones also provoke God's "
     "anger (Numbers 12:8-10). God's anger is likewise upon those who hinder the "
     "preaching of the gospel (1 Thessalonians 2:15-16). It is aroused by "
     "immorality, suppression of the truth, unrepentance, despising God's words, "
     "mocking at His people, covetousness, envy, and the like (Colossians 3:5-6; "
     "2 Chronicles 36:15-16). False worship and apostasy can directly provoke God "
     "to anger.")

# ---- Possess your soul in anger ----
section_heading("Possess your soul in anger")
lead("Unjustified and uncontrolled anger has led many persons into greater sins, "
     "even acts of violence. Cain grew hot with great anger and slew Abel "
     "(Genesis 4:5, 8). Esau wanted to kill Jacob, who had received the blessing "
     "of their father (Genesis 27:41-45). Saul, in his rage, hurled spears at "
     "David and at Jonathan (1 Samuel 18:11; 19:10; 20:33).")
bullet("Those in attendance at the synagogue in Nazareth, aroused to anger by "
       "Jesus' preaching, endeavoured to hurl Him from the brow of a mountain "
       "(Luke 4:28-29).")
bullet("Angered religious leaders rushed upon Stephen with one accord and stoned "
       "him to death (Acts 7:54-60).")
lead("Anger, even when justified, if not controlled, may be dangerous, producing "
     "bad results.")
bullet("Simeon and Levi had reason to be indignant at Shechem for violating their "
       "sister Dinah. But the wanton slaughter of the Shechemites was an excessive "
       "penalty to inflict. Hence their father Jacob denounced their uncontrolled "
       "anger, cursing it (Genesis 34:1-31; 49:5-7).")
bullet("When under heavy provocation, a person should control his anger. The "
       "complaint and rebelliousness of the Israelites provoked Moses, the meekest "
       "man on the land, to an uncontrolled act of anger, in which he failed to "
       "sanctify God, and for which he was punished (Numbers 12:3; 20:10-12; "
       "Psalm 106:32-33).")
bullet("Fruits of anger are classified along with other detestable works of the "
       "flesh (Galatians 5:19-21).")
bullet("Angry talk is to be kept out of the congregation. Men representing the "
       "congregation in prayer should be free from feelings of anger and ill will "
       "(1 Timothy 2:8).")
bullet("Christians are commanded to be slow about wrath, being told that man's "
       "wrath does not work out God's righteousness (James 1:19-20).")
bullet("Christians are further counselled to yield place to the wrath, and to "
       "leave vengeance to God (Romans 12:19).")
bullet("A man cannot be used as an overseer of God if he is prone to wrath "
       "(Titus 1:7).")
bullet("While a person may on occasion be justifiably angry, he should not let it "
       "become sin to him by harbouring it or maintaining a provoked state. He "
       "should not let the sun set with him in such a condition, for he will "
       "thereby allow place for the Devil to take advantage of him "
       "(Ephesians 4:26-27).")
bullet("Christians are advised to take proper steps to make peace in the Lord, in "
       "the way provided (Leviticus 19:17-18; Matthew 5:23-24; Luke 17:3-4).")
bullet("The Scriptures counsel that we should watch our associations in this "
       "regard, not keeping company with anyone given to anger or fits of rage, "
       "and so avoiding a snare for our souls (Proverbs 22:24-25).")

# ---- Causes of human anger ----
section_heading("Causes of human anger")
lead("In the Bible, anger is caused by human sin, wounded pride, injustice, and "
     "frustration, as well as by a righteous concern for God's honour. Scripture "
     "distinguishes between destructive, sinful, human anger and justified, holy "
     "anger driven by love and zeal for truth. Common causes of human anger are:")
bullet2("Pride and ego. ", "Refusing to accept disrespect, slights, or loss of "
        "control (Proverbs 21:24).")
bullet2("Unmet desires. ", "Wanting something so badly that fighting or "
        "bitterness results when it is withheld (James 4:1-3).")
bullet2("Foolishness and impatience. ", "Speaking or reacting too quickly, "
        "without listening or thinking (Proverbs 14:29; James 1:19-20).")
bullet2("Envy and jealousy. ", "Resenting someone else's success, possessions, or "
        "favour, as seen in Cain's anger toward Abel (Genesis 4:5-6).")
bullet2("Hurt and bitterness. ", "Harbouring deep resentment or unaddressed "
        "emotional pain that turns into an explosive or slow-burning temper "
        "(Ephesians 4:31).")

# ---- Causes of righteous anger ----
section_heading("Causes of righteous anger")
bullet2("Injustice and oppression. ", "Seeing the vulnerable mistreated or abused "
        "against God's moral law (Exodus 2:11-12).")
bullet2("Disrespect for God. ", "Witnessing the corruption of holy things or "
        "blatant rebellion against God's commands, as seen when Jesus cleared the "
        "temple (John 2:14-17).")
bullet2("Hardness of heart. ", "Grieving over human spiritual blindness and "
        "refusal to show mercy (Mark 3:5).")

# ---- Christ's example ----
section_heading("Anger: Christ's example")
lead("The records of the life of Christ do not recount one occasion where He was "
     "guilty of uncontrolled anger, or where He allowed the lawlessness, "
     "rebelliousness, and harassment of the enemies of God to upset His spirit and "
     "cause Him to reflect such a thing towards His followers or others.")
lead("When He drove out those who were defiling God's temple, as well as violating "
     "the laws of Moses by making God's house a house of merchandise, it was "
     "through no uncontrolled, unjustified fit of anger. Rather, the Scriptures "
     "show that it was properly directed zeal for the house of God "
     "(John 2:13-17).")

# ---- Damaging effects ----
section_heading("Damaging effects of anger")
lead("Not only does anger have an adverse effect upon our spiritual health, but it "
     "produces profound effects on the physical organism.")
bullet("It can cause a rise in blood pressure, arterial changes, respiratory "
       "troubles, liver upsets, changes in the secretion of gall, and effects on "
       "the pancreas.")
bullet("Anger and rage, as strong emotions, have been listed by physicians as "
       "contributing to, aggravating, or even causing such ailments as asthma, eye "
       "afflictions, skin diseases, hives, ulcers, and dental and digestive "
       "troubles.")
bullet("Rage and fury can upset thinking processes so that one cannot form "
       "logical conclusions or pass sound judgement.")
bullet("The aftermath of a fit of rage is often a period of extreme mental "
       "depression.")
lead("It is therefore wisdom, not only in a religious sense but in a physical "
     "sense, to keep anger under control and to pursue peace and love "
     "(Proverbs 14:29-30; Romans 14:19; 1 Peter 3:11). The Christian should strive "
     "to control his spirit, avoiding the destructive emotion of anger as the "
     "world is drawing to an end (Proverbs 14:29; Ecclesiastes 7:9).")

# ---- Conclusion ----
section_heading("Conclusion")
lead("According to thanatologists, those who study people who are dying:")
bullet("Most sick people deny the truth that they are ill (“No, not me”).")
bullet("When they accept its reality, they often feel enormous anger (“Why "
       "me?”).")
bullet("Depression sets in after the bargaining fails to work.")
bullet("Support from loved ones and professional counselling can help the sick to "
       "accept their approaching death with courage and die in peace.")
lead("Indeed, life begins and ends with anger, but let us choose peace over anger "
     "(Romans 14:19).")

# ---- Tit-bits ----
section_heading("Tit-bits about anger")
bullet("Anger should not lead a Christian into sin (Psalm 37:8).")
bullet("Those prone to anger should not lead the church in prayer "
       "(1 Timothy 2:8).")
bullet("The discretion of a man makes him slow to anger (Proverbs 19:11).")
bullet("He who is slow to anger is better than the mighty (Proverbs 16:32).")
bullet("Wrath is cruel and anger is a torrent (Proverbs 27:4).")
bullet("A quick-tempered man acts foolishly (Proverbs 14:17).")
bullet("Whoever is angry with his brother without a cause shall be in danger of "
       "the judgement (Matthew 5:22).")
bullet("Avoid angry people (Proverbs 22:24).")
bullet("Put off all these: anger and wrath (Colossians 3:8).")
bullet("Solomon advises us as Christians not to harbour anger "
       "(Ecclesiastes 7:9).")

# ================= SAVE =================
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Anger - draft.docx")
try:
    doc.save(out)
except PermissionError:
    raise SystemExit("Cannot write '%s'. It is open in Word. Close it and run again." % out)
print("Saved:", out)
print("Paragraphs:", len(doc.paragraphs))
