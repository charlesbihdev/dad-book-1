# -*- coding: utf-8 -*-
"""Build the 'Christian Holiness' chapter as a Word doc matching the style of
'Introduction - final.docx' and the other chapters: Garamond, A5, 13pt justified.
Six sheets, TYPED with the scripture in Twi (Akan) and the headings in English.
Per the project rule the Twi scripture is REPLACED with the English NIV, keeping
the English headings and structure. The opening paragraph is Dad's handwritten
definition on sheet 1. Quotations are italic. References in plain parentheses,
hyphens in verse ranges, no em dashes."""
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Garamond"
BODY_SZ = Pt(13)

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = BODY_SZ
rpr = normal.element.get_or_add_rPr()
rf = rpr.get_or_add_rFonts()
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

sec = doc.sections[0]
sec.page_width = Mm(148)
sec.page_height = Mm(210)
sec.top_margin = Mm(14)
sec.bottom_margin = Mm(15)
sec.left_margin = Mm(15)
sec.right_margin = Mm(13)

def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p)

def _set_run(r, bold=False, italic=False, size=None):
    r.font.name = FONT
    rpr = r._r.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    r.font.bold = bold
    r.font.italic = italic
    if size: r.font.size = size

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper()); _set_run(r, bold=True, size=Pt(22))
    return p

def subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text); _set_run(r, bold=True, italic=True, size=Pt(14))
    return p

def point(num, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.keep_with_next = True
    r = p.add_run("%d.  " % num); _set_run(r, bold=True)
    r = p.add_run(text); _set_run(r, bold=True)
    return p

def letter(mark, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-6)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(1)
    p.keep_with_next = True
    r = p.add_run(mark + "  "); _set_run(r, bold=True)
    r = p.add_run(text); _set_run(r, bold=True)
    return p

def lead(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); _set_run(r)
    return p

def leadq(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(1)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r)
    return p

def scripture(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(10)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text); _set_run(r, italic=True, size=Pt(12))
    return p

LQ = "“"; RQ = "”"; LS = "‘"; RS = "’"
def q(text):
    return LQ + text + RQ

# ================= CONTENT =================

title("Christian Holiness")
subtitle("Let us be holy, as God is holy")

lead("Christian holiness refers to a state where a Christian turns away from sin "
     "and conforms his life to God's standards rather than the world's pattern. "
     "For a Christian, approaching a holy God requires dealing with impurity, so "
     "that people can safely live in relationship with Him.")

# ---- 1 ----
point(1, "Does God require holiness from us?")
leadq("God told the Israelites to be holy as Himself (Leviticus 11:44):")
scripture(q("I am the Lord your God; consecrate yourselves and be holy, because I "
            "am holy. Do not make yourselves unclean by any creature that moves "
            "along the ground."))
leadq("In the same way God asks His people to partake in His holiness "
      "(1 Peter 1:14-16):")
scripture(q("As obedient children, do not conform to the evil desires you had when "
            "you lived in ignorance. But just as he who called you is holy, so be "
            "holy in all you do; for it is written: " + LS + "Be holy, because I am "
            "holy." + RS))
leadq("God warned the Jews not to do what the Egyptians and Canaanites were doing "
      "(Leviticus 18:1-4):")
scripture(q("The Lord said to Moses, " + LS + "Speak to the Israelites and say to "
            "them: I am the Lord your God. You must not do as they do in Egypt, "
            "where you used to live, and you must not do as they do in the land of "
            "Canaan, where I am bringing you. Do not follow their practices. You "
            "must obey my laws and be careful to follow my decrees. I am the Lord "
            "your God." + RS))
leadq("The standard for righteousness is not set by following the multitude "
      "(Exodus 23:2):")
scripture(q("Do not follow the crowd in doing wrong. When you give testimony in a "
            "lawsuit, do not pervert justice by siding with the crowd."))
leadq("Christ put it this way (Matthew 7:13-14):")
scripture(q("Enter through the narrow gate. For wide is the gate and broad is the "
            "road that leads to destruction, and many enter through it. But small "
            "is the gate and narrow the road that leads to life, and only a few "
            "find it."))

# ---- 2 ----
point(2, "What prevents us from being holy?")
leadq("Ungodly practices that destroy many (Galatians 5:19-21):")
scripture(q("The acts of the flesh are obvious: sexual immorality, impurity and "
            "debauchery; idolatry and witchcraft; hatred, discord, jealousy, fits "
            "of rage, selfish ambition, dissensions, factions and envy; "
            "drunkenness, orgies, and the like. I warn you, as I did before, that "
            "those who live like this will not inherit the kingdom of God."))
leadq("There are sexual sins, adultery, fornication, uncleanness, lasciviousness "
      "(Romans 1:26-32):")
scripture(q("Because of this, God gave them over to shameful lusts. Even their "
            "women exchanged natural sexual relations for unnatural ones. In the "
            "same way the men also abandoned natural relations with women and were "
            "inflamed with lust for one another. Men committed shameful acts with "
            "other men, and received in themselves the due penalty for their "
            "error. Furthermore, just as they did not think it worthwhile to "
            "retain the knowledge of God, so God gave them over to a depraved "
            "mind, so that they do what ought not to be done. They have become "
            "filled with every kind of wickedness, evil, greed and depravity. They "
            "are full of envy, murder, strife, deceit and malice. They are "
            "gossips, slanderers, God-haters, insolent, arrogant and boastful; "
            "they invent ways of doing evil; they disobey their parents. They have "
            "no understanding, no fidelity, no love, no mercy. Although they know "
            "God's righteous decree that those who do such things deserve death, "
            "they not only continue to do these very things but also approve of "
            "those who practice them."))
leadq("There are the sins of idolatry and witchcraft (Psalm 115:3-8):")
scripture(q("Our God is in heaven; he does whatever pleases him. But their idols "
            "are silver and gold, made by human hands. They have mouths, but "
            "cannot speak, eyes, but cannot see. They have ears, but cannot hear, "
            "noses, but cannot smell. They have hands, but cannot feel, feet, but "
            "cannot walk, nor can they utter a sound with their throats. Those who "
            "make them will be like them, and so will all who trust in them."))
leadq("Let us crucify sin (Colossians 3:5-10):")
scripture(q("Put to death, therefore, whatever belongs to your earthly nature: "
            "sexual immorality, impurity, lust, evil desires and greed, which is "
            "idolatry. Because of these, the wrath of God is coming. You used to "
            "walk in these ways, in the life you once lived. But now you must also "
            "rid yourselves of all such things as these: anger, rage, malice, "
            "slander, and filthy language from your lips. Do not lie to each "
            "other, since you have taken off your old self with its practices and "
            "have put on the new self, which is being renewed in knowledge in the "
            "image of its Creator."))
lead("There are sins of attitude, hatred, variance, emulations, wrath, jealousy.")
leadq("There are the sins of false teaching and division, strife, seditions, "
      "heresies (1 Timothy 4:1-2):")
scripture(q("The Spirit clearly says that in later times some will abandon the "
            "faith and follow deceiving spirits and things taught by demons. Such "
            "teachings come through hypocritical liars, whose consciences have "
            "been seared as with a hot iron."))
leadq("There are sins of drunkenness and revellings (1 Peter 4:3):")
scripture(q("For you have spent enough time in the past doing what pagans choose "
            "to do, living in debauchery, lust, drunkenness, orgies, carousing and "
            "detestable idolatry."))
leadq("Sins of worldly living (Ephesians 2:1-3):")
scripture(q("As for you, you were dead in your transgressions and sins, in which "
            "you used to live when you followed the ways of this world and of the "
            "ruler of the kingdom of the air, the spirit who is now at work in "
            "those who are disobedient. All of us also lived among them at one "
            "time, gratifying the cravings of our flesh and following its desires "
            "and thoughts. Like the rest, we were by nature deserving of wrath."))
leadq("John put it this way (1 John 2:15-17):")
scripture(q("Do not love the world or anything in the world. If anyone loves the "
            "world, love for the Father is not in them. For everything in the "
            "world, the lust of the flesh, the lust of the eyes, and the pride of "
            "life, comes not from the Father but from the world. The world and its "
            "desires pass away, but whoever does the will of God lives forever."))
leadq("Whoever befriends the world becomes an enemy of God (James 4:4-5):")
scripture(q("You adulterous people, don't you know that friendship with the world "
            "means enmity against God? Therefore, anyone who chooses to be a "
            "friend of the world becomes an enemy of God. Or do you think "
            "Scripture says without reason that he jealously longs for the spirit "
            "he has caused to dwell in us?"))

# ---- 3 ----
point(3, "Is holiness possible for Christians?")
leadq("Peter said it is possible for Christians to be holy (1 Peter 2:21-25):")
scripture(q("To this you were called, because Christ suffered for you, leaving you "
            "an example, that you should follow in his steps. " + LS + "He "
            "committed no sin, and no deceit was found in his mouth." + RS + " When "
            "they hurled their insults at him, he did not retaliate; when he "
            "suffered, he made no threats. Instead, he entrusted himself to him "
            "who judges justly. " + LS + "He himself bore our sins" + RS + " in his "
            "body on the cross, so that we might die to sins and live for "
            "righteousness; " + LS + "by his wounds you have been healed." + RS +
            " For " + LS + "you were like sheep going astray," + RS + " but now you "
            "have returned to the Shepherd and Overseer of your souls."))
leadq("We are to bear the fruits of the Holy Spirit (Galatians 5:22-26):")
scripture(q("But the fruit of the Spirit is love, joy, peace, forbearance, "
            "kindness, goodness, faithfulness, gentleness and self-control. "
            "Against such things there is no law. Those who belong to Christ Jesus "
            "have crucified the flesh with its passions and desires. Since we live "
            "by the Spirit, let us keep in step with the Spirit. Let us not become "
            "conceited, provoking and envying each other."))

# ---- 4 ----
point(4, "What role can the church play?")
letter("(a)", "The church is to uphold holiness (Revelation 2:5):")
scripture(q("Consider how far you have fallen! Repent and do the things you did at "
            "first. If you do not repent, I will come to you and remove your "
            "lampstand from its place."))
leadq("When the church fails in morality, she is rejected (Matthew 5:13):")
scripture(q("You are the salt of the earth. But if the salt loses its saltiness, "
            "how can it be made salty again? It is no longer good for anything, "
            "except to be thrown out and trampled underfoot."))
leadq("(Mark 9:50):")
scripture(q("Salt is good, but if it loses its saltiness, how can you make it "
            "salty again? Have salt among yourselves, and be at peace with each "
            "other."))
letter("(b)", "The church is duty-bound not to condone sin (Romans 1:32):")
scripture(q("Although they know God's righteous decree that those who do such "
            "things deserve death, they not only continue to do these very things "
            "but also approve of those who practice them."))
leadq("Malachi put it this way (Malachi 2:17):")
scripture(q("You have wearied the Lord with your words. " + LS + "How have we "
            "wearied him?" + RS + " you ask. By saying, " + LS + "All who do evil "
            "are good in the eyes of the Lord, and he is pleased with them" + RS +
            " or " + LS + "Where is the God of justice?" + RS))
letter("(c)", "The church is to rebuke sin (1 Samuel 13:13):")
scripture(q(LS + "You have done a foolish thing," + RS + " Samuel said. " + LS +
            "You have not kept the command the Lord your God gave you; if you had, "
            "he would have established your kingdom over Israel for all time." + RS))
leadq("Paul was clear on this (1 Corinthians 5:1-2):")
scripture(q("It is actually reported that there is sexual immorality among you, "
            "and of a kind that even pagans do not tolerate: A man is sleeping "
            "with his father's wife. And you are proud! Shouldn't you rather have "
            "gone into mourning and have put out of your fellowship the man who "
            "has been doing this?"))
leadq("(1 Corinthians 5:6-8):")
scripture(q("Your boasting is not good. Don't you know that a little yeast leavens "
            "the whole batch of dough? Get rid of the old yeast, so that you may "
            "be a new unleavened batch, as you really are. For Christ, our "
            "Passover lamb, has been sacrificed. Therefore let us keep the "
            "Festival, not with the old bread leavened with malice and wickedness, "
            "but with the unleavened bread of sincerity and truth."))
letter("(d)", "The church may apply the tool of disfellowship "
        "(1 Corinthians 5:9-13):")
scripture(q("I wrote to you in my letter not to associate with sexually immoral "
            "people, not at all meaning the people of this world who are immoral, "
            "or the greedy and swindlers, or idolaters. In that case you would "
            "have to leave this world. But now I am writing to you that you must "
            "not associate with anyone who claims to be a brother or sister but is "
            "sexually immoral or greedy, an idolater or slanderer, a drunkard or "
            "swindler. Do not even eat with such people. What business is it of "
            "mine to judge those outside the church? Are you not to judge those "
            "inside? God will judge those outside. " + LS + "Expel the wicked "
            "person from among you." + RS))
leadq("Paul again put it this way (2 Thessalonians 3:6):")
scripture(q("In the name of the Lord Jesus Christ, we command you, brothers and "
            "sisters, to keep away from every believer who is idle and disruptive "
            "and does not live according to the teaching you received from us."))
letter("(e)", "The church should teach the benefits of holiness "
        "(2 Corinthians 7:1):")
scripture(q("Therefore, since we have these promises, dear friends, let us purify "
            "ourselves from everything that contaminates body and spirit, "
            "perfecting holiness out of reverence for God."))

leadq("Members are to imitate Christ (Ephesians 4:20-24):")
scripture(q("That, however, is not the way of life you learned when you heard "
            "about Christ and were taught in him in accordance with the truth that "
            "is in Jesus. You were taught, with regard to your former way of life, "
            "to put off your old self, which is being corrupted by its deceitful "
            "desires; to be made new in the attitude of your minds; and to put on "
            "the new self, created to be like God in true righteousness and "
            "holiness."))
leadq("Without peace and holiness no one can see God (Hebrews 12:14-16):")
scripture(q("Make every effort to live in peace with everyone and to be holy; "
            "without holiness no one will see the Lord. See to it that no one "
            "falls short of the grace of God and that no bitter root grows up to "
            "cause trouble and defile many. See that no one is sexually immoral, "
            "or is godless like Esau, who for a single meal sold his inheritance "
            "rights as the oldest son."))
leadq("Let us be holy, as God is holy (1 Peter 1:14-16):")
scripture(q("As obedient children, do not conform to the evil desires you had when "
            "you lived in ignorance. But just as he who called you is holy, so be "
            "holy in all you do; for it is written: " + LS + "Be holy, because I am "
            "holy." + RS))

# ================= SAVE =================
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Christian Holiness - draft.docx")
try:
    doc.save(out)
except PermissionError:
    raise SystemExit("Cannot write '%s'. It is open in Word. Close it and run again." % out)
print("Saved:", out)
print("Paragraphs:", len(doc.paragraphs))
