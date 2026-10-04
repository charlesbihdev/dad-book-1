# -*- coding: utf-8 -*-
"""Build the Conscience chapter (handwritten sheets 1-9) as a Word doc matching
the style of 'Introduction - final.docx' and the other chapters: Garamond, A5,
13pt justified. References in plain parentheses, hyphens in verse ranges, no em
dashes. Quoted verses are given in the NIV, in italics. The chapter is organised
as an Introduction, Steps I-VI, and 'The three pictures of conscience'."""
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Garamond"
BODY_SZ = Pt(13)

doc = Document()

# ---- base style ----
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = BODY_SZ
rpr = normal.element.get_or_add_rPr()
rf = rpr.get_or_add_rFonts()
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

# ---- page = A5, margins from the first book ----
sec = doc.sections[0]
sec.page_width = Mm(148)
sec.page_height = Mm(210)
sec.top_margin = Mm(14)
sec.bottom_margin = Mm(15)
sec.left_margin = Mm(15)
sec.right_margin = Mm(13)

# ---- footer page number ----
def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p)

# ---- helpers ----
def _set_run(r, bold=False, italic=False, size=None):
    r.font.name = FONT
    rpr = r._r.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    r.font.bold = bold
    r.font.italic = italic
    if size: r.font.size = size

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(text.upper())
    _set_run(r, bold=True, size=Pt(22))
    return p

def section_heading(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(15))
    return p

def step_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(13))
    return p

def qheading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(13))
    return p

def point_label(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(13))
    return p

def lead(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text); _set_run(r)
    return p

def bullet(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("•  "); _set_run(r)
    r = p.add_run(text); _set_run(r)
    return p

def scripture(text):
    """A quoted verse (NIV), italic and indented."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text); _set_run(r, italic=True, size=Pt(12))
    return p

# ================= CONTENT =================

title("Conscience")

# ---- Introduction ----
section_heading("Introduction")
qheading("What is conscience?")
lead("The word conscience is translated from the Greek word “syneidesis”, "
     "which is drawn from “syn” (with) and “eidesis” (knowledge), "
     "and thus means co-knowledge, or knowledge with oneself. Therefore, "
     "conscience is a capacity to look at oneself and render judgement about "
     "oneself. The apostle Paul expresses the operation of his conscience in this "
     "manner:")
scripture("“I speak the truth in Christ, I am not lying; my conscience "
          "confirms it through the Holy Spirit” (Romans 9:1).")

# ---- Step I ----
step_heading("Step I: Man is born with conscience")
lead("Conscience is inherent in man, having been made part of him by God. It is "
     "an inward realization or sense of right and wrong that excuses or accuses a "
     "person for his actions. Hence conscience judges. It can also be trained by "
     "the thoughts and acts, convictions and rules that are implanted in a "
     "person's mind by study and experience. Based on these things, it makes a "
     "comparison with the course of action being taken. Then it sounds a warning "
     "when the rules and the course conflict, unless the conscience is "
     "“seared”, made unfeeling by continued violations of its warnings. "
     "Conscience can be a moral safety device, because it imparts pleasure when "
     "the action is good and inflicts pain when the action taken is bad.")
lead("Man has had conscience from creation. Adam and Eve manifested this as soon "
     "as they broke God's laws and hid themselves (Genesis 3:7). It can also be "
     "seen that conscience has not been wiped out even among non-Christians "
     "(Romans 2:14-15). This is because all mankind descended from Adam and Eve, "
     "in whom conscience is inherent. Many laws of the nations are in harmony with "
     "a Christian's conscience, yet such nations and lawmakers may not have been "
     "influenced by Christianity at all. The laws were made according to the "
     "leadings of their own consciences. All persons have the faculty of "
     "conscience, and it is to this that the life, course, and preaching of "
     "Christianity appeal (2 Corinthians 4:2; Acts 2:37):")
scripture("“Rather, we have renounced secret and shameful ways; we do not use "
          "deception, nor do we distort the word of God. On the contrary, by "
          "setting forth the truth plainly we commend ourselves to everyone's "
          "conscience in the sight of God” (2 Corinthians 4:2).")
scripture("“When the people heard this, they were cut to the heart and said "
          "to Peter and the other apostles, ‘Brothers, what shall we "
          "do?’” (Acts 2:37).")

# ---- Step II ----
step_heading("Step II: Conscience needs training")
lead("Conscience must be enlightened. If not, it can mislead. It is an unsafe "
     "guide if it has not been trained in right standards, according to the "
     "truth. Its development can be wrongly influenced by local environment, "
     "customs, worship, and habits. It might judge matters as being right or "
     "wrong by those incorrect standards or values. An example of this is shown "
     "when Christ indicated that men will even kill God's servants, thinking that "
     "they were doing Him a service (John 16:2).")
scripture("According to Hamilton: “Conscience is universal because there is a "
          "certain characteristic innate in the mind which enables a person who "
          "has reached the age of reasoning ability to make judgement as to the "
          "rightness or wrongness of any course of action which may be presented "
          "to the mind. Faced with a particular course of action, the mind "
          "instinctively reacts with the corresponding judgement: ‘I ought to "
          "do the right.’”")
lead("Saul, later Paul the apostle, went out with murderous intent against "
     "Christ's disciples, believing he was zealously serving God "
     "(Galatians 1:13-16). The Jews were seriously misled into fighting against "
     "God because of a lack of appreciation of God's word (Romans 10:1-3; "
     "Acts 4:1-3; Acts 5:39-40). Only a conscience properly trained by God's word "
     "can correctly assess and set matters of life thoroughly straight "
     "(Hebrews 4:12). A Christian must have a stable, right standard, God's "
     "standard.")

# ---- Step III ----
step_heading("Step III: Good conscience")
lead("One must approach God with a cleansed conscience (Hebrews 10:22). A "
     "Christian must constantly strive for an honest conscience in all things "
     "(Hebrews 13:18). When Paul stated:")
scripture("“So I strive always to keep my conscience clear before God and "
          "man” (Acts 24:16),")
lead("he meant that he continually steered and corrected his course of life "
     "according to God's words and Christ's teachings, for he knew that in the "
     "final analysis God, and not his own conscience, would be his ultimate judge "
     "(1 Corinthians 4:4). Following a Bible-trained conscience may result in "
     "persecution, but Peter the apostle comforted us that:")
scripture("“For it is commendable if someone bears up under the pain of "
          "unjust suffering because they are conscious of God” (1 Peter 2:19).")
lead("A Christian must hold a good conscience in the face of opposition "
     "(1 Peter 3:16). The law, with its animal sacrifices, could not so perfect a "
     "person as regards his conscience that he could consider himself free from "
     "guilt. However, through the application of Christ's ransom to those having "
     "faith, a person's conscience can be cleansed (Hebrews 9:9, 14). Peter "
     "indicates that those who receive salvation have to have this good, clean, "
     "right conscience (1 Peter 3:21).")

# ---- Step IV ----
step_heading("Step IV: Bad conscience")
lead("The conscience can be so abused that it no longer is clean and sensitive. "
     "When that happens, it cannot sound out warnings or give safe guidance "
     "(Titus 1:15). As a result, man's conduct is then controlled by fear of "
     "exposure and punishment rather than by a good conscience (Romans 13:4-5). "
     "Paul's reference to a conscience that is marked as with a branding iron "
     "indicates that it would be like seared flesh that is covered over with scar "
     "tissue and void of nerve endings, and therefore without sense of feeling "
     "(1 Timothy 4:1-2). Persons with such a conscience cannot sense right or "
     "wrong. They do not appreciate the freedom God grants them and, rebelling, "
     "become a slave to a bad conscience. It is easy to defile one's conscience. "
     "A Christian's aim should be as shown in Acts 23:1:")
scripture("“My brothers, I have fulfilled my duty to God in all good "
          "conscience to this day” (Acts 23:1).")

# ---- Step V ----
step_heading("Step V: Consider the consciences of others")
lead("In order to make proper evaluations, a conscience must be fully and "
     "accurately trained in God's word. An untrained conscience may be weak. As a "
     "result, it may be easily and unwisely suppressed, or the person may become "
     "offended by the actions or words of others, even in instances where no "
     "wrongdoing may exist. Paul gave examples of this in connection with eating, "
     "drinking, and the judging of certain days as above others (Romans 14:1-23; "
     "1 Corinthians 8:1-13). The Christian with knowledge and whose conscience is "
     "trained is commanded to give consideration and allowance to the one with a "
     "weak conscience, not using all his freedom or insisting on all his personal "
     "rights or always doing just as he pleases (Romans 15:1). One who wounds the "
     "weak conscience of a fellow Christian is sinning against Christ "
     "(1 Corinthians 8:12). On the other hand, Paul implies that while he would "
     "not want to do something by which the weak brother would be offended, "
     "thereby causing the brother to judge Paul, the weak one should likewise "
     "consider his brother, striving for maturity by getting more knowledge and "
     "training so that his conscience will not be easily offended, causing him to "
     "view others wrongly (1 Corinthians 10:29-30; Romans 14:10).")

# ---- Step VI ----
step_heading("Step VI: Demonstrations of conscience")
lead("In the Bible, the conscience is demonstrated as an inner moral witness that "
     "accuses or excuses behaviour.")
bullet("Adam and Eve felt immediate guilt, shame, and fear, hiding behind the "
       "trees after eating the forbidden fruit (Genesis 3:7-8).")
bullet("David's heart smote him, his conscience was pricked, just for cutting off "
       "a small piece of King Saul's robe (1 Samuel 24:5).")
bullet("Joseph's brothers spoke, years after selling him into slavery, saying "
       "their current troubles were punishment for their guilty deeds "
       "(Genesis 42:21). Also, after realizing that, with their father's death, "
       "Joseph might take revenge and deal with them ruthlessly for their inhuman "
       "treatment of him, they went to him to plead for mercy (Genesis 50:15-18).")
bullet("The men trying to stone the adulterous woman were convicted by their own "
       "conscience and left one by one (John 8:9).")
bullet("Gentiles who do not have the written law still have God's moral law "
       "written on their hearts (Romans 2:14-15).")
bullet("We should let our conscience lead us to give cheerfully "
       "(2 Corinthians 9:6-7).")
bullet("The Bible admonishes us to examine ourselves, whether we are drifting out "
       "of the faith (2 Corinthians 13:5).")
bullet("In the approach to and observance of the Lord's Supper, the application "
       "of a clean conscience is paramount (1 Corinthians 11:27-29).")
bullet("Judas returned his betrayal money to the Jewish authorities after his "
       "conscience pricked him (Matthew 27:3-5).")

# ---- The three pictures ----
section_heading("The three pictures of conscience")
point_label("(i) The court room (Romans 2:14-15)")
bullet("Here, the human conscience acts as an internal courtroom where a person "
       "stands trial.")
bullet("Thoughts, words, and deeds are brought forward as evidence.")
bullet("The internal witness accuses the person, and the conscience either "
       "defends them (excusing) or prosecutes them (accusing with guilt).")
point_label("(ii) The umpire")
bullet("The conscience acts like a sports official making split-second calls on "
       "behaviour.")
bullet("It reviews an action and sounds an alarm, often described as a red light "
       "on the dashboard of the heart, when a moral boundary is crossed.")
point_label("(iii) The window (Titus 1:15)")
bullet("The condition of the conscience affects how a person perceives reality.")
bullet("A clean conscience provides clear spiritual light, while a corrupted or "
       "“dirty” conscience distorts perception, making everything look "
       "defiled or unclean to the observer.")
lead("In short, the conscience needs constant cleansing (Hebrews 9:14). An evil "
     "or guilty conscience cannot be fixed purely by human effort or ritual, but "
     "is cleansed internally through faith and the blood of Christ.")

# ================= SAVE =================
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Conscience - draft.docx")
try:
    doc.save(out)
except PermissionError:
    raise SystemExit("Cannot write '%s'. It is open in Word. Close it and run again." % out)
print("Saved:", out)
print("Paragraphs:", len(doc.paragraphs))
