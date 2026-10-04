# -*- coding: utf-8 -*-
"""Build the Dream chapter as a Word doc matching the style of
'Introduction - final.docx' and the other chapters: Garamond, A5, 13pt justified.
This chapter was typed with the scripture in Twi (Akan) and English headers; per
the project rule the Twi scripture is REPLACED with the English NIV (taken from an
authoritative NIV source), keeping the same structure and headers. Quotations are
italic. References in plain parentheses, hyphens in verse ranges, no em dashes.
Sections are renumbered 1-6 (the sheets skipped from 3 to 5)."""
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Garamond"
BODY_SZ = Pt(13)

doc = Document()

# ---- base style ----
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = BODY_SZ
rpr = normal.element.get_or_add_rPr()
rf = rpr.get_or_add_rFonts()
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

# ---- page = A5 ----
sec = doc.sections[0]
sec.page_width = Mm(148)
sec.page_height = Mm(210)
sec.top_margin = Mm(14)
sec.bottom_margin = Mm(15)
sec.left_margin = Mm(15)
sec.right_margin = Mm(13)

# ---- footer page number ----
def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p)

# ---- helpers ----
def _set_run(r, bold=False, italic=False, size=None):
    r.font.name = FONT
    rpr = r._r.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    r.font.bold = bold
    r.font.italic = italic
    if size: r.font.size = size

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(text.upper())
    _set_run(r, bold=True, size=Pt(22))
    return p

def qheading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(13))
    return p

def label(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(13))
    return p

def lead(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(2)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r)
    return p

def bullet(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("•  "); _set_run(r)
    r = p.add_run(text); _set_run(r)
    return p

def scripture(text):
    """A quoted verse (NIV), italic and indented."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text); _set_run(r, italic=True, size=Pt(12))
    return p

LQ = "“"; RQ = "”"; LS = "‘"; RS = "’"
def q(text):
    return LQ + text + RQ

# ================= CONTENT =================

title("Dreams")

# ---- 1 ----
qheading("1. Give us your introduction.")
lead("In the Old Testament, God communicated to man through dreams, visions, "
     "supernatural happenings, and written words of inspiration.")
lead("But in these last days, God communicates to us through Christ "
     "(Hebrews 1:1-2):")
scripture(q("In the past God spoke to our ancestors through the prophets at many "
            "times and in various ways, but in these last days he has spoken to us "
            "by his Son, whom he appointed heir of all things, and through whom "
            "also he made the universe."))
lead("God is the interpreter of dreams (Genesis 40:8):")
scripture(q("We both had dreams,") + " they answered, " + q("but there is no one "
          "to interpret them.") + " Then Joseph said to them, " + q("Do not "
          "interpretations belong to God? Tell me your dreams."))
lead("There is also the example of Daniel (Daniel 2:19-23):")
scripture(q("During the night the mystery was revealed to Daniel in a vision. "
            "Then Daniel praised the God of heaven and said: " + LS + "Praise be "
            "to the name of God for ever and ever; wisdom and power are his. He "
            "changes times and seasons; he deposes kings and raises up others. He "
            "gives wisdom to the wise and knowledge to the discerning. He reveals "
            "deep and hidden things; he knows what lies in darkness, and light "
            "dwells with him. I thank and praise you, God of my ancestors: You "
            "have given me wisdom and power, you have made known to me what we "
            "asked of you, you have made known to us the dream of the king." + RS))

# ---- 2 ----
qheading("2. What is a dream?")
lead("A dream refers to the thoughts or mental images a person has while asleep. "
     "The Bible describes three types of dream.")
label("Dreams from God (Numbers 12:6):")
scripture(q("Listen to my words: When there is a prophet among you, I, the Lord, "
            "reveal myself to them in visions, I speak to them in dreams."))
label("Natural dreams.")
label("False dreams (Jeremiah 29:8-9):")
scripture(q("Yes, this is what the Lord Almighty, the God of Israel, says: " + LS +
            "Do not let the prophets and diviners among you deceive you. Do not "
            "listen to the dreams you encourage them to have. They are prophesying "
            "lies to you in my name. I have not sent them," + RS + " declares the "
            "Lord."))

# ---- 3 ----
qheading("3. What do you mean by a dream from God?")
lead("Dreams from God were received both by God's servants and by those who were "
     "not His servants.")
label("(a) Some dreams gave warnings, like that of the king of Gerar, Abimelech "
      "(Genesis 20:2-3):")
scripture(q("Abraham said of his wife Sarah, " + LS + "She is my sister." + RS +
            " Then Abimelek king of Gerar sent for Sarah and took her. But God "
            "came to Abimelek in a dream one night and said to him, " + LS + "You "
            "are as good as dead because of the woman you have taken; she is a "
            "married woman." + RS))
label("(b) Other dreams gave guidance, like that of the Magi (Matthew 2:11-12):")
scripture(q("On coming to the house, they saw the child with his mother Mary, and "
            "they bowed down and worshiped him. Then they opened their treasures "
            "and presented him with gifts of gold, frankincense and myrrh. And "
            "having been warned in a dream not to go back to Herod, they returned "
            "to their country by another route."))
lead("Joseph, through a dream from God, took Mary as his wife (Matthew 1:20):")
scripture(q("But after he had considered this, an angel of the Lord appeared to "
            "him in a dream and said, " + LS + "Joseph son of David, do not be "
            "afraid to take Mary home as your wife, because what is conceived in "
            "her is from the Holy Spirit." + RS))
lead("Joseph was also warned to flee to Egypt to save the baby Christ "
     "(Matthew 2:13):")
scripture(q("When they had gone, an angel of the Lord appeared to Joseph in a "
            "dream. " + LS + "Get up," + RS + " he said, " + LS + "take the child "
            "and his mother and escape to Egypt. Stay there until I tell you, for "
            "Herod is going to search for the child to kill him." + RS))
label("(c) Assurance of divine favour was also given through dreams.")
lead("Abraham was promised the land of Canaan (Genesis 15:12-16):")
scripture(q("As the sun was setting, Abram fell into a deep sleep, and a thick "
            "and dreadful darkness came over him. Then the Lord said to him, " + LS +
            "Know for certain that for four hundred years your descendants will be "
            "strangers in a country not their own and that they will be enslaved "
            "and mistreated there. But I will punish the nation they serve as "
            "slaves, and afterward they will come out with great possessions. You, "
            "however, will go to your ancestors in peace and be buried at a good "
            "old age. In the fourth generation your descendants will come back "
            "here, for the sin of the Amorites has not yet reached its full "
            "measure." + RS))
lead("Joseph's future blessings were revealed through dreams (Genesis 37:5-10):")
scripture(q("Joseph had a dream, and when he told it to his brothers, they hated "
            "him all the more. He said to them, " + LS + "Listen to this dream I "
            "had: We were binding sheaves of grain out in the field when suddenly "
            "my sheaf rose and stood upright, while your sheaves gathered around "
            "mine and bowed down to it." + RS + " His brothers said to him, " + LS +
            "Do you intend to reign over us? Will you actually rule us?" + RS +
            " And they hated him all the more because of his dream and what he had "
            "said. Then he had another dream, and he told it to his brothers. " +
            LS + "Listen," + RS + " he said, " + LS + "I had another dream, and "
            "this time the sun and moon and eleven stars were bowing down to "
            "me." + RS + " When he told his father as well as his brothers, his "
            "father rebuked him and said, " + LS + "What is this dream you had? "
            "Will your mother and I and your brothers actually come and bow down "
            "to the ground before you?" + RS))
label("(d) Some dreams given to those who were not worshippers of God were "
      "prophetic (Genesis 40:5-8):")
scripture(q("Each of the two men, the cupbearer and the baker of the king of "
            "Egypt, who were being held in prison, had a dream the same night, and "
            "each dream had a meaning of its own. When Joseph came to them the "
            "next morning, he saw that they were dejected. So he asked Pharaoh's "
            "officials who were in custody with him in his master's house, " + LS +
            "Why do you look so sad today?" + RS + " " + LS + "We both had "
            "dreams," + RS + " they answered, " + LS + "but there is no one to "
            "interpret them." + RS + " Then Joseph said to them, " + LS + "Do not "
            "interpretations belong to God? Tell me your dreams." + RS))
lead("Pharaoh had a prophetic dream of seven lean cows swallowing seven fat cows "
     "(Genesis 41:25-30):")
scripture(q("Then Joseph said to Pharaoh, " + LS + "The dreams of Pharaoh are one "
            "and the same. God has revealed to Pharaoh what he is about to do. The "
            "seven good cows are seven years, and the seven good heads of grain are "
            "seven years; it is one and the same dream. The seven lean, ugly cows "
            "that came up afterward are seven years, and so are the seven worthless "
            "heads of grain scorched by the east wind: They are seven years of "
            "famine. It is just as I said to Pharaoh: God has shown Pharaoh what he "
            "is about to do. Seven years of great abundance are coming throughout "
            "the land of Egypt, but seven years of famine will follow them." + RS))
lead("This prophetic dream was a way of saving Jacob and his household, sending "
     "them to Egypt to fulfil God's promise to Abraham (Genesis 45:5-8):")
scripture(q("And now, do not be distressed and do not be angry with yourselves for "
            "selling me here, because it was to save lives that God sent me ahead "
            "of you. For two years now there has been famine in the land, and for "
            "the next five years there will be no plowing and reaping. But God sent "
            "me ahead of you to preserve for you a remnant on earth and to save "
            "your lives by a great deliverance. So then, it was not you who sent me "
            "here, but God. He made me father to Pharaoh, lord of his entire "
            "household and ruler of all Egypt."))
lead("Nebuchadnezzar's first dream was about God's kingdom and the kingdoms of "
     "Babylon, Medo-Persia, Greece, and Rome (Daniel 2:44-45):")
scripture(q("In the time of those kings, the God of heaven will set up a kingdom "
            "that will never be destroyed, nor will it be left to another people. "
            "It will crush all those kingdoms and bring them to an end, but it "
            "will itself endure forever. This is the meaning of the vision of the "
            "rock cut out of a mountain, but not by human hands, a rock that broke "
            "the iron, the bronze, the clay, the silver and the gold to pieces."))
lead("In a second dream, Nebuchadnezzar was to live like a beast for seven years "
     "(Daniel 4:24-27):")
scripture(q("You will be driven away from people and will live with the wild "
            "animals; you will eat grass like the ox and be drenched with the dew "
            "of heaven. Seven times will pass by for you until you acknowledge "
            "that the Most High is sovereign over all kingdoms on earth and gives "
            "them to anyone he wishes. The command to leave the stump of the tree "
            "with its roots means that your kingdom will be restored to you when "
            "you acknowledge that Heaven rules. Therefore, Your Majesty, be "
            "pleased to accept my advice: Renounce your sins by doing what is "
            "right, and your wickedness by being kind to the oppressed."))
lead("Pilate's wife was warned in a dream (Matthew 27:19):")
scripture(q("While Pilate was sitting on the judge's seat, his wife sent him this "
            "message: " + LS + "Don't have anything to do with that innocent man, "
            "for I have suffered a great deal today in a dream because of "
            "him." + RS))

# ---- 4 ----
qheading("4. Explain to us what you mean by natural dreams.")
lead("Natural dreams may be stimulated by certain thoughts or emotions, or by "
     "daily activities such as anxiety, one's physical condition, or one's "
     "occupation (Ecclesiastes 5:3):")
scripture(q("A dream comes when there are many cares, and many words mark the "
            "speech of a fool."))
bullet("These dreams are of no effect.")
bullet("A hungry person may dream of eating.")
bullet("A poor person may dream of having money (Isaiah 29:8):")
scripture(q("As when a hungry person dreams of eating, but awakens hungry still; "
            "as when a thirsty person dreams of drinking, but awakens faint and "
            "thirsty still."))

# ---- 5 ----
qheading("5. What are false dreams?")
lead("False dreamers claim to receive a message from God about people or about "
     "themselves (Jeremiah 23:25-32):")
scripture(q("I have heard what the prophets say who prophesy lies in my name. "
            "They say, " + LS + "I had a dream! I had a dream!" + RS + " How long "
            "will this continue in the hearts of these lying prophets, who "
            "prophesy the delusions of their own minds? They think the dreams they "
            "tell one another will make my people forget my name, just as their "
            "ancestors forgot my name through Baal worship. Let the prophet who has "
            "a dream recount the dream, but let the one who has my word speak it "
            "faithfully. For what has straw to do with grain? declares the Lord. "
            "Is not my word like fire, declares the Lord, and like a hammer that "
            "breaks a rock in pieces? Therefore, declares the Lord, I am against "
            "the prophets who steal from one another words supposedly from me. "
            "Yes, declares the Lord, I am against the prophets who wag their own "
            "tongues and yet declare, " + LS + "The Lord declares." + RS + " "
            "Indeed, I am against those who prophesy false dreams, declares the "
            "Lord. They tell them and lead my people astray with their reckless "
            "lies, yet I did not send or appoint them. They do not benefit these "
            "people in the least, declares the Lord."))
lead("False prophets also use dreams to give false hope (Zechariah 10:2):")
scripture(q("The idols speak deceitfully, diviners see visions that lie; they "
            "tell dreams that are false, they give comfort in vain."))

# ---- 6 ----
qheading("6. Does God speak to us in dreams with the coming of Christianity?")
lead("No. God now speaks to us through Christ (Hebrews 1:1-2):")
scripture(q("In the past God spoke to our ancestors through the prophets at many "
            "times and in various ways, but in these last days he has spoken to us "
            "by his Son, whom he appointed heir of all things, and through whom "
            "also he made the universe."))
lead("The Holy Spirit guides us into all truth (John 16:13):")
scripture(q("But when he, the Spirit of truth, comes, he will guide you into all "
            "the truth. He will not speak on his own; he will speak only what he "
            "hears, and he will tell you what is yet to come."))
lead("All Scripture is given by inspiration (2 Timothy 3:16-17):")
scripture(q("All Scripture is God-breathed and is useful for teaching, rebuking, "
            "correcting and training in righteousness, so that the servant of God "
            "may be thoroughly equipped for every good work."))
lead("The truth has, once and for all time, been revealed (2 Peter 1:3):")
scripture(q("His divine power has given us everything we need for a godly life "
            "through our knowledge of him who called us by his own glory and "
            "goodness."))
lead("Jude warns us (Jude 1:3-4):")
scripture(q("Dear friends, although I was very eager to write to you about the "
            "salvation we share, I felt compelled to write and urge you to contend "
            "for the faith that was once for all entrusted to God's holy people. "
            "For certain individuals whose condemnation was written about long ago "
            "have secretly slipped in among you. They are ungodly people, who "
            "pervert the grace of our God into a license for immorality and deny "
            "Jesus Christ our only Sovereign and Lord."))
lead("Paul said he was astonished when people accept another gospel "
     "(Galatians 1:6-7):")
scripture(q("I am astonished that you are so quickly deserting the one who called "
            "you to live in the grace of Christ and are turning to a different "
            "gospel, which is really no gospel at all."))
lead("Too much dreaming is useless (Ecclesiastes 5:7):")
scripture(q("Much dreaming and many words are meaningless. Therefore fear God."))

# ================= SAVE =================
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Dream - draft.docx")
try:
    doc.save(out)
except PermissionError:
    raise SystemExit("Cannot write '%s'. It is open in Word. Close it and run again." % out)
print("Saved:", out)
print("Paragraphs:", len(doc.paragraphs))
