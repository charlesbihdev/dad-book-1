# -*- coding: utf-8 -*-
"""Build the 'False Prophets' chapter as a Word doc matching the style of
'Introduction - final.docx' and the other chapters: Garamond, A5, 13pt justified.
Three sheets, TYPED with the scripture in Twi (Akan) and the headings in English.
Per the project rule the Twi scripture is REPLACED with the English NIV, keeping
the English headings and structure. Two sections: the characteristics of false
prophets (1-5) and what we should do with them (1-4). Quotations are italic.
References in plain parentheses, hyphens in verse ranges, no em dashes."""
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Garamond"
BODY_SZ = Pt(13)

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = BODY_SZ
rpr = normal.element.get_or_add_rPr()
rf = rpr.get_or_add_rFonts()
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

sec = doc.sections[0]
sec.page_width = Mm(148)
sec.page_height = Mm(210)
sec.top_margin = Mm(14)
sec.bottom_margin = Mm(15)
sec.left_margin = Mm(15)
sec.right_margin = Mm(13)

def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p)

def _set_run(r, bold=False, italic=False, size=None):
    r.font.name = FONT
    rpr = r._r.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    r.font.bold = bold
    r.font.italic = italic
    if size: r.font.size = size

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text.upper()); _set_run(r, bold=True, size=Pt(22))
    return p

def heading(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(14))
    return p

def point(num, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.keep_with_next = True
    r = p.add_run("%d.  " % num); _set_run(r, bold=True)
    r = p.add_run(text); _set_run(r, bold=True)
    return p

def bullet(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("•  "); _set_run(r)
    r = p.add_run(text); _set_run(r)
    return p

def subq(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(1)
    p.keep_with_next = True
    r = p.add_run("•  "); _set_run(r)
    r = p.add_run(text); _set_run(r)
    return p

def leadq(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(1)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r)
    return p

def scripture(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(10)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text); _set_run(r, italic=True, size=Pt(12))
    return p

LQ = "“"; RQ = "”"; LS = "‘"; RS = "’"
def q(text):
    return LQ + text + RQ

# ================= CONTENT =================

title("False Prophets")

heading("The Characteristics of False Prophets")

# ---- 1 ----
point(1, "They may make predictions that do not come true (Deuteronomy 18:21-22):")
scripture(q("You may say to yourselves, " + LS + "How can we know when a message "
            "has not been spoken by the Lord?" + RS + " If what a prophet proclaims "
            "in the name of the Lord does not take place or come true, that is a "
            "message the Lord has not spoken. That prophet has spoken "
            "presumptuously, so do not be alarmed."))

# ---- 2 ----
point(2, "They may perform miraculous signs and wonders (2 Thessalonians 2:9-11):")
scripture(q("The coming of the lawless one will be in accordance with how Satan "
            "works. He will use all sorts of displays of power through signs and "
            "wonders that serve the lie, and all the ways that wickedness deceives "
            "those who are perishing. They perish because they refused to love the "
            "truth and so be saved. For this reason God sends them a powerful "
            "delusion so that they will believe the lie."))
subq("Moses and the Egyptian sorcerers (Exodus 7:11-12):")
scripture(q("Pharaoh then summoned wise men and sorcerers, and the Egyptian "
            "magicians also did the same things by their secret arts: Each one "
            "threw down his staff and it became a snake. But Aaron's staff "
            "swallowed up their staffs."))
subq("Simon the sorcerer (Acts 8:9-11):")
scripture(q("Now for some time a man named Simon had practiced sorcery in the city "
            "and amazed all the people of Samaria. He boasted that he was someone "
            "great, and all the people, both high and low, gave him their "
            "attention and exclaimed, " + LS + "This man is rightly called the "
            "Great Power of God." + RS + " They followed him because he had amazed "
            "them for a long time with his sorcery."))

# ---- 3 ----
point(3, "They claim to be Christ (Matthew 24:4-5):")
scripture(q("Jesus answered: " + LS + "Watch out that no one deceives you. For "
            "many will come in my name, claiming, " + LQ + "I am the Messiah," + RQ +
            " and will deceive many." + RS))
leadq("(Matthew 24:24-26):")
scripture(q("For false messiahs and false prophets will appear and perform great "
            "signs and wonders to deceive, if possible, even the elect. See, I "
            "have told you ahead of time. So if anyone tells you, " + LS + "There "
            "he is, out in the wilderness," + RS + " do not go out; or, " + LS +
            "Here he is, in the inner rooms," + RS + " do not believe it."))

# ---- 4 ----
point(4, "They may have unbiblical lifestyles.")
bullet("They live immorally.")
bullet("They defy authorities.")
bullet("They do whatever their instinct tells them.")
bullet("They deceive people for money.")
bullet("They care only for themselves.")
subq("They bear no fruit (Jude 8):")
scripture(q("In the very same way, on the strength of their dreams these ungodly "
            "people pollute their own bodies, reject authority and heap abuse on "
            "celestial beings."))
leadq("(Jude 10-12):")
scripture(q("Yet these people slander whatever they do not understand, and the "
            "very things they do understand by instinct, as irrational animals do, "
            "will destroy them. Woe to them! They have taken the way of Cain; they "
            "have rushed for profit into Balaam's error; they have been destroyed "
            "in Korah's rebellion. These people are blemishes at your love feasts, "
            "eating with you without the slightest qualm, shepherds who feed only "
            "themselves. They are clouds without rain, blown along by the wind; "
            "autumn trees, without fruit and uprooted, twice dead."))
leadq("(Jude 15-19):")
scripture(q("to judge everyone, and to convict all of them of all the ungodly "
            "acts they have committed in their ungodliness, and of all the defiant "
            "words ungodly sinners have spoken against him. These people are "
            "grumblers and faultfinders; they follow their own evil desires; they "
            "boast about themselves and flatter others for their own advantage. "
            "But, dear friends, remember what the apostles of our Lord Jesus "
            "Christ foretold. They said to you, " + LS + "In the last times there "
            "will be scoffers who will follow their own ungodly desires." + RS +
            " These are the people who divide you, who follow mere natural "
            "instincts and do not have the Spirit."))

# ---- 5 ----
point(5, "Their teachings lead people away from Christ (Luke 6:40):")
scripture(q("The student is not above the teacher, but everyone who is fully "
            "trained will be like their teacher."))

# ================= SECTION 2 =================
heading("What Should We Do With False Prophets?")

# ---- 1 ----
point(1, "We should not be afraid of them (Deuteronomy 18:22):")
scripture(q("If what a prophet proclaims in the name of the Lord does not take "
            "place or come true, that is a message the Lord has not spoken. That "
            "prophet has spoken presumptuously, so do not be alarmed."))

# ---- 2 ----
point(2, "Let us avoid them (Romans 16:17-18):")
scripture(q("I urge you, brothers and sisters, to watch out for those who cause "
            "divisions and put obstacles in your way that are contrary to the "
            "teaching you have learned. Keep away from them. For such people are "
            "not serving our Lord Christ, but their own appetites. By smooth talk "
            "and flattery they deceive the minds of naive people."))

# ---- 3 ----
point(3, "Let us expose them (Ephesians 5:11):")
scripture(q("Have nothing to do with the fruitless deeds of darkness, but rather "
            "expose them."))

# ---- 4 ----
point(4, "Spoil their business, as Demetrius said (Acts 19:23-28):")
scripture(q("About that time there arose a great disturbance about the Way. A "
            "silversmith named Demetrius, who made silver shrines of Artemis, "
            "brought in a lot of business for the craftsmen there. He called them "
            "together, along with the workers in related trades, and said: " + LS +
            "You know, my friends, that we receive a good income from this "
            "business. And you see and hear how this fellow Paul has convinced and "
            "led astray large numbers of people here in Ephesus and in practically "
            "the whole province of Asia. He says that gods made by human hands are "
            "no gods at all. There is danger not only that our trade will lose its "
            "good name, but also that the temple of the great goddess Artemis will "
            "be discredited; and the goddess herself, who is worshiped throughout "
            "the province of Asia and the world, will be robbed of her divine "
            "majesty." + RS + " When they heard this, they were furious and began "
            "shouting: " + LS + "Great is Artemis of the Ephesians!" + RS))
subq("Simon the sorcerer (Acts 8:9-13):")
scripture(q("Now for some time a man named Simon had practiced sorcery in the city "
            "and amazed all the people of Samaria. He boasted that he was someone "
            "great, and all the people, both high and low, gave him their "
            "attention and exclaimed, " + LS + "This man is rightly called the "
            "Great Power of God." + RS + " They followed him because he had amazed "
            "them for a long time with his sorcery. But when they believed Philip "
            "as he proclaimed the good news of the kingdom of God and the name of "
            "Jesus Christ, they were baptized, both men and women. Simon himself "
            "believed and was baptized. And he followed Philip everywhere, "
            "astonished by the great signs and miracles he saw."))

# ================= SAVE =================
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "False Prophets - draft.docx")
try:
    doc.save(out)
except PermissionError:
    raise SystemExit("Cannot write '%s'. It is open in Word. Close it and run again." % out)
print("Saved:", out)
print("Paragraphs:", len(doc.paragraphs))
