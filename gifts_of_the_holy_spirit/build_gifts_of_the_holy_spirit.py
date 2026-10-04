# -*- coding: utf-8 -*-
"""Build the 'Gifts of the Holy Spirit' chapter as a Word doc matching the style
of 'Introduction - final.docx' and the other chapters: Garamond, A5, 13pt
justified.

UNLIKE the other chapters, this one is NOT Dad's own manuscript: the folder holds
two photos of a printed reference-book article headed 'GIFTS FROM GOD' (pp. 937-
938), with Dad's handwritten topic labels on top. Per Charles's instruction it is
treated as an ordinary source page of the book and transcribed here.

The article is prose with scripture CITATIONS (no quoted verse text), so there is
no Twi-to-NIV conversion. Editorial handling, all noted in the corrections file:
 - abbreviated citations expanded to full book names;
 - end-of-sentence em-dash citations normalised to the book's parenthetical style
   (no em dashes anywhere);
 - the source's internal cross-references ('see FAITH; HEALING', 'see PROPHECY;
   PROPHET ...') removed, as they point to other articles in that encyclopedia;
 - the source's wording 'Jehovah' and 'C.E.' kept as printed and flagged for Dad.
"""
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Garamond"
BODY_SZ = Pt(13)

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = BODY_SZ
rpr = normal.element.get_or_add_rPr()
rf = rpr.get_or_add_rFonts()
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

sec = doc.sections[0]
sec.page_width = Mm(148)
sec.page_height = Mm(210)
sec.top_margin = Mm(14)
sec.bottom_margin = Mm(15)
sec.left_margin = Mm(15)
sec.right_margin = Mm(13)

def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p)

def _set_run(r, bold=False, italic=False, size=None):
    r.font.name = FONT
    rpr = r._r.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    r.font.bold = bold
    r.font.italic = italic
    if size: r.font.size = size

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text.upper()); _set_run(r, bold=True, size=Pt(22))
    return p

def para(body, lead=None, lead_italic=False):
    """A justified paragraph with an optional bold (optionally italic) lead-in run."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    if lead:
        r = p.add_run(lead + "  "); _set_run(r, bold=True, italic=lead_italic)
    r = p.add_run(body); _set_run(r)
    return p

LQ = "“"; RQ = "”"
def q(text):
    return LQ + text + RQ

# ================= CONTENT =================

title("Gifts of the Holy Spirit")

para("In the first century C.E. miraculous gifts attended the baptism with holy "
     "spirit. These served as signs that God was no longer using the Jewish "
     "congregation in his service but that his approval rested on the Christian "
     "congregation established by his Son. (Hebrews 2:2-4) On the day of "
     "Pentecost, miraculous gifts accompanied the outpouring of the holy spirit, "
     "and in each case mentioned thereafter in the Scriptures where the miraculous "
     "gifts of the spirit were transmitted, at least one of the 12 apostles or "
     "Paul, who was directly chosen by Jesus, was present. (Acts 2:1, 4, 14; "
     "8:9-20; 10:44-46; 19:6) Evidently, with the death of the apostles, the "
     "transmittal of the gifts of the spirit ended, and the miraculous gifts of "
     "the spirit ceased altogether as those who had received these gifts passed "
     "off the earthly scene.",
     lead="Gifts of the Spirit.")

para("Performing apparently miraculous works would not in itself prove divine "
     "authorization, nor would the inability of God's servants to perform miracles "
     "with the help of God's spirit cast doubt on the fact that they were being "
     "used by him. (Matthew 7:21-23) Not every first-century Christian could "
     "perform powerful works, heal, speak in tongues, and translate. Paul, and "
     "doubtless some others, had by God's undeserved kindness been granted a "
     "number of these gifts of the spirit. However, these miraculous gifts marked "
     "the infancy of the Christian congregation and were foretold to cease. In "
     "fact, even Jesus indicated that his followers would be identified, not by "
     "their performance of powerful works, but by their love for one another. "
     "(1 Corinthians 12:29, 30; 13:2, 8-13; John 13:35)")

para("Paul enumerates nine different manifestations or operations of the spirit: "
     "(1) speech of wisdom, (2) speech of knowledge, (3) faith, (4) gifts of "
     "healings, (5) powerful works, (6) prophesying, (7) discernment of inspired "
     "utterances, (8) different tongues, and (9) interpretation of tongues. All "
     "these gifts of the spirit served a beneficial purpose that not only "
     "contributed to the numerical growth of the congregation but also resulted "
     "in its spiritual upbuilding. (1 Corinthians 12:7-11; 14:24-26)")

para("Although wisdom can be acquired through study, application, and experience, "
     "the " + q("speech of wisdom") + " here mentioned apparently was a miraculous "
     "ability to apply knowledge in a successful way to solve problems arising in "
     "the congregation. (1 Corinthians 12:8) It was " + q("according to the wisdom "
     "given him") + " that Paul wrote letters that became part of God's inspired "
     "Word. (2 Peter 3:15, 16) This gift also appears to have been manifest in the "
     "individual's ability to make a defense that opposers were unable to resist "
     "or to dispute. (Acts 6:9, 10)",
     lead=q("Speech of wisdom."), lead_italic=True)

para("All in the first-century Christian congregation had basic knowledge "
     "concerning Jehovah and his Son as well as God's will and his requirements "
     "for life. Therefore, " + q("speech of knowledge") + " was something above "
     "and beyond the knowledge shared by Christians in general; it was miraculous "
     "knowledge. Likewise " + q("faith") + " as a gift of the spirit was evidently "
     "a miraculous faith that helped the individual to overcome mountainlike "
     "obstacles that would otherwise hinder service to God. "
     "(1 Corinthians 12:8, 9; 13:2)",
     lead=q("Speech of knowledge") + " and " + q("faith."), lead_italic=True)

para("The gift of healing was manifest in the ability to cure diseases "
     "completely, regardless of the nature of the affliction. (Acts 5:15, 16; "
     "9:33, 34; 28:8, 9) Prior to Pentecost, healing had been done by Jesus and "
     "his disciples. Whereas some persons healed did manifest obvious faith, the "
     "afflicted one was not required to make an expression of faith in order to be "
     "cured. (Compare John 5:5-9, 13.) Jesus, on one occasion, attributed his "
     "disciples' inability to cure an epileptic, not to the lack of faith of the "
     "one seeking a cure for his son, but to the little faith of his disciples. "
     "(Matthew 17:14-16, 18-20) Not once do the Scriptures cite an instance where "
     "Jesus or his apostles were unable to heal others on account of the lack of "
     "faith of those seeking a cure. Furthermore, instead of using the gift of "
     "healing in curing Timothy of his stomach trouble or attributing his frequent "
     "cases of sickness to his lack of faith, the apostle Paul recommended that "
     "Timothy use a little wine for the sake of his stomach. (1 Timothy 5:23)",
     lead=q("Healings."), lead_italic=True)

para("Powerful works included raising dead persons, expelling demons, and even "
     "striking opposers with blindness. (1 Corinthians 12:10) The manifestation "
     "of such powerful works resulted in adding believers to the congregation. "
     "(Acts 9:40, 42; 13:8-12; 19:11, 12, 20)",
     lead=q("Powerful works."), lead_italic=True)

para("Prophesying was a greater gift than speaking in tongues, as it built up the "
     "congregation. Moreover, unbelievers were helped thereby to recognize that "
     "God was really among the Christians. (1 Corinthians 14:3-5, 24, 25) All in "
     "the Christian congregation spoke about the fulfillment of the prophecies "
     "recorded in God's Word. (Acts 2:17, 18) However, the particular ones having "
     "the miraculous gift of prophesying were able to foretell future events, as "
     "did Agabus. (Acts 11:27, 28)",
     lead=q("Prophesying."), lead_italic=True)

para("Discernment of inspired utterances evidently involved the ability to "
     "discern whether an inspired expression originated with God or not. "
     "(1 Corinthians 12:10) This gift would prevent its possessor from being "
     "deceived and turned away from the truth and would protect the congregation "
     "from false prophets. (1 John 4:1; compare 2 Corinthians 11:3, 4)",
     lead=q("Discernment of inspired utterances."), lead_italic=True)

para("The miraculous gift of tongues attended the outpouring of God's spirit at "
     "Pentecost, 33 C.E. The approximately 120 disciples assembled in an upper "
     "room (possibly near the temple) were thereby enabled to speak about " +
     q("the magnificent things of God") + " in the native tongues of the Jews and "
     "proselytes who had come to Jerusalem from faraway places for the observance "
     "of the festival. This fulfillment of Joel's prophecy proved that God was "
     "using the new Christian congregation and no longer the Jewish congregation. "
     "In order to receive the free gift of the holy spirit, the Jews and "
     "proselytes had to repent and be baptized in Jesus' name. (Acts 1:13-15; "
     "2:1-47)",
     lead=q("Tongues."), lead_italic=True)

para("The gift of tongues proved very helpful to first-century Christians in "
     "preaching to those who spoke other languages. It was actually a sign to "
     "unbelievers. However, Paul, in writing to the Christian congregation at "
     "Corinth, directed that when meeting together, not all should speak in "
     "tongues, as strangers and unbelievers entering and not understanding would "
     "conclude that they were mad. He also recommended that the speaking in "
     "tongues " + q("be limited to two or three at the most, and in turns.") +
     " However, if no one could translate, then the person speaking in a tongue "
     "was to remain silent in the congregation, speaking to himself and to God. "
     "(1 Corinthians 14:22-33) If no translating took place, his speaking in a "
     "tongue would not result in upbuilding others, for no one would listen to his "
     "speech because it would be meaningless to those unable to understand it. "
     "(1 Corinthians 14:2, 4)")

para("If the person speaking in a tongue was unable to translate, then he did not "
     "understand what he himself was saying nor would others who were not familiar "
     "with that tongue, or language. Hence, Paul encouraged those having the gift "
     "of tongues to pray that they might also translate and thereby edify all "
     "listeners. From the foregoing, it can readily be seen why Paul, under "
     "inspiration, ranked speaking in tongues as a lesser gift and pointed out "
     "that in a congregation he would rather speak five words with his mind "
     "(understanding) than 10,000 words in a tongue. (1 Corinthians 14:11, 13-19)")

para("The gift of interpretation of tongues was manifest in a person's being able "
     "to translate a language unknown to the one having this gift. "
     "(1 Corinthians 12:10) This gift really enhanced the gift of speaking in "
     "tongues, since the entire congregation would be built up by hearing the "
     "translation. (1 Corinthians 14:5)",
     lead=q("Interpretation of tongues."), lead_italic=True)

# ================= SAVE =================
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Gifts of the Holy Spirit - draft.docx")
try:
    doc.save(out)
except PermissionError:
    raise SystemExit("Cannot write '%s'. It is open in Word. Close it and run again." % out)
print("Saved:", out)
print("Paragraphs:", len(doc.paragraphs))
