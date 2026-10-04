# -*- coding: utf-8 -*-
"""Build the 'Imitators' chapter as a Word doc matching the style of
'Introduction - final.docx' and the other chapters: Garamond, A5, 13pt justified.
This chapter is MIXED: sheets 1 and 3 were TYPED with the scripture in Twi (Akan)
and are REPLACED with the English NIV (keeping the English headings); sheet 2 and
the lower part of sheet 1 are HANDWRITTEN English with inline references.
Quotations are italic. References in plain parentheses, hyphens in verse ranges,
no em dashes."""
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Garamond"
BODY_SZ = Pt(13)

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = BODY_SZ
rpr = normal.element.get_or_add_rPr()
rf = rpr.get_or_add_rFonts()
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

sec = doc.sections[0]
sec.page_width = Mm(148)
sec.page_height = Mm(210)
sec.top_margin = Mm(14)
sec.bottom_margin = Mm(15)
sec.left_margin = Mm(15)
sec.right_margin = Mm(13)

def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p)

def _set_run(r, bold=False, italic=False, size=None):
    r.font.name = FONT
    rpr = r._r.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    r.font.bold = bold
    r.font.italic = italic
    if size: r.font.size = size

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper()); _set_run(r, bold=True, size=Pt(22))
    return p

def subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text); _set_run(r, bold=True, italic=True, size=Pt(14))
    return p

def point(num, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.keep_with_next = True
    r = p.add_run("%d.  " % num); _set_run(r, bold=True)
    r = p.add_run(text); _set_run(r, bold=True)
    return p

def label(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(13))
    return p

def lead(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); _set_run(r)
    return p

def hang(marker, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(8)
    p.paragraph_format.first_line_indent = Mm(-8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(marker); _set_run(r)
    r = p.add_run(text); _set_run(r)
    return p

def bullet(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("•  "); _set_run(r)
    r = p.add_run(text); _set_run(r)
    return p

def leadq(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(1)
    p.keep_with_next = True
    r = p.add_run("•  "); _set_run(r)
    r = p.add_run(text); _set_run(r)
    return p

def scripture(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(10)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text); _set_run(r, italic=True, size=Pt(12))
    return p

LQ = "“"; RQ = "”"; LS = "‘"; RS = "’"
def q(text):
    return LQ + text + RQ

# ================= CONTENT =================

title("Imitators")
subtitle("Let us be imitators of Christ")

label("Text: John 13:14-15")
scripture(q("Now that I, your Lord and Teacher, have washed your feet, you also "
            "should wash one another's feet. I have set you an example that you "
            "should do as I have done for you."))

label("Introduction")
hang("(i)   ", "To imitate means to copy the behaviour of a person, or to take "
     "someone's behaviour (the good) as an example, or to be like someone.")
hang("(ii)  ", "Another word close to the word imitation is " + q("pattern") +
     ", which refers to an excellent example.")

label("Uses of examples")

# ---- 1 ----
point(1, "Examples give warnings to believers of the gospel.")
leadq("(1 Corinthians 10:5-6):")
scripture(q("Nevertheless, God was not pleased with most of them; their bodies "
            "were scattered in the wilderness. Now these things occurred as "
            "examples to keep us from setting our hearts on evil things as they "
            "did."))
leadq("(2 Peter 2:4-7):")
scripture(q("For if God did not spare angels when they sinned, but sent them to "
            "hell, putting them in chains of darkness to be held for judgment; if "
            "he did not spare the ancient world when he brought the flood on its "
            "ungodly people, but protected Noah, a preacher of righteousness, and "
            "seven others; if he condemned the cities of Sodom and Gomorrah by "
            "burning them to ashes, and made them an example of what is going to "
            "happen to the ungodly; and if he rescued Lot, a righteous man, who "
            "was distressed by the depraved conduct of the lawless."))

# ---- 2 ----
point(2, "Examples give models of life for believers' attention.")
bullet("They set a pattern of suffering (1 Peter 2:21-23).")
bullet("They teach us not to follow bad examples (Hebrews 4:11).")
bullet("They teach us good deeds to be imitated from Christ (John 13:14-15).")
bullet("As believers, let us imitate the sinless life of Christ, who could ask, " +
       q("Can any of you prove me guilty of sin?") + " (John 8:46).")
bullet("By imitating Christ, we should have a forgiving spirit, as He prayed, " +
       q("Father, forgive them") + " (Luke 23:34).")
bullet("We should imitate Christ by equipping ourselves with God's word "
       "(Deuteronomy 8:3; Matthew 4:4).")
bullet("Jesus was the doer of all that He taught.")
bullet("Jesus finished His work and, at His death, committed Himself to the "
       "Father, saying, " + q("Father, into your hands I commit my spirit") +
       " (Luke 23:46).")
bullet("Christ was humble (Matthew 11:29-30).")
bullet("We learn from the example of Christ's unshakeable faith "
       "(Philippians 1:27).")
bullet("Examples show the manner in which sincere believers responded to the will "
       "of God.")

# ---- 3 ----
point(3, "Examples are used to edify the members of the church, teaching them how "
         "to:")
leadq("Live faithful lives (Revelation 2:13):")
scripture(q("I know where you live, where Satan has his throne. Yet you remain "
            "true to my name. You did not renounce your faith in me, not even in "
            "the days of Antipas, my faithful witness, who was put to death in "
            "your city, where Satan lives."))
leadq("Lead zealous lives (Acts 8:4):")
scripture(q("Those who had been scattered preached the word wherever they went."))
leadq("Lead generous lives (2 Corinthians 8:1-5):")
scripture(q("And now, brothers and sisters, we want you to know about the grace "
            "that God has given the Macedonian churches. In the midst of a very "
            "severe trial, their overflowing joy and their extreme poverty welled "
            "up in rich generosity. For I testify that they gave as much as they "
            "were able, and even beyond their ability. Entirely on their own, they "
            "urgently pleaded with us for the privilege of sharing in this service "
            "to the Lord's people. And they exceeded our expectations: They gave "
            "themselves first of all to the Lord, and then by the will of God also "
            "to us."))
bullet("Examples motivate, stimulate and encourage us to respond to God's "
       "principles.")

# ---- 4 ----
point(4, "Examples help believers to clarify and apply God's commands.")

# ---- conclusion ----
label("Conclusion")
lead("Every Christian should be able to declare, as Paul did (1 Corinthians 11:1):")
scripture(q("Follow my example, as I follow the example of Christ."))

# ================= SAVE =================
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Imitators - draft.docx")
try:
    doc.save(out)
except PermissionError:
    raise SystemExit("Cannot write '%s'. It is open in Word. Close it and run again." % out)
print("Saved:", out)
print("Paragraphs:", len(doc.paragraphs))
