# -*- coding: utf-8 -*-
"""Build the Lie chapter (handwritten sheets 1-4) as a Word doc matching the
style of 'Introduction - final.docx' and the other chapters: Garamond, A5, 13pt
justified. References in plain parentheses, hyphens in verse ranges, no em dashes.
The chapter is a single teaching essay under the question 'What is a lie?'."""
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Garamond"
BODY_SZ = Pt(13)

doc = Document()

# ---- base style ----
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = BODY_SZ
rpr = normal.element.get_or_add_rPr()
rf = rpr.get_or_add_rFonts()
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

# ---- page = A5, margins from the first book ----
sec = doc.sections[0]
sec.page_width = Mm(148)
sec.page_height = Mm(210)
sec.top_margin = Mm(14)
sec.bottom_margin = Mm(15)
sec.left_margin = Mm(15)
sec.right_margin = Mm(13)

# ---- footer page number ----
def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p)

# ---- helpers ----
def _set_run(r, bold=False, italic=False, size=None):
    r.font.name = FONT
    rpr = r._r.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    r.font.bold = bold
    r.font.italic = italic
    if size: r.font.size = size

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(text.upper())
    _set_run(r, bold=True, size=Pt(22))
    return p

def qheading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(13))
    return p

def lead(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text); _set_run(r)
    return p

def bullet(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("•  "); _set_run(r)
    r = p.add_run(text); _set_run(r)
    return p

# ================= CONTENT =================

title("Lie")

qheading("What is a lie?")
lead("A lie is the opposite of truth. Lying generally involves saying something "
     "false to a person who is entitled to know the truth, and doing so with the "
     "intent to deceive or to injure him or another person.")

bullet("A lie need not always be verbal. It can also be expressed in action; that "
       "is, a person may be living a lie.")
bullet("Basically, a lie may refer to something worthless, vain, or valueless "
       "(Psalm 12:2; Zechariah 10:2).")
bullet("The father, or originator, of lying is Satan the Devil (John 8:44). His "
       "lie, conveyed by means of a serpent to the first woman, Eve, ultimately "
       "brought death to her and her husband, Adam (Genesis 3:1-5, 16-19).")
bullet("That first lie was rooted in selfishness and wrong desire. It was "
       "designed to divert the love and obedience of the human pair to the liar, "
       "who had presented himself as an angel of light, a benefactor "
       "(2 Corinthians 11:13-14).")
bullet("All other malicious lies uttered since that time have likewise been a "
       "reflection of selfishness and wrong desire.")
bullet("People have told lies to escape deserved punishment, to profit at the "
       "expense of others, and to gain or maintain certain advantages, material "
       "reward, or the praise of men.")
bullet("Especially serious have been the religious lies, as they have endangered "
       "the future life of persons deceived by them (Matthew 23:15). The exchange "
       "of God's truth for “the lie”, the falsehood of idolatry, can cause a "
       "person to become a practicer of what is degrading and vile "
       "(Romans 1:24-32).")
bullet("The case of the religious leaders of Judaism in the time of Jesus' "
       "earthly ministry shows what can happen when one abandons the truth.")
bullet("They schemed to have Jesus put to death (Matthew 27:1-2).")
bullet("Then, when He was resurrected, they bribed the soldiers who had guarded "
       "the tomb so they would conceal the truth and spread a lie about the "
       "disappearance of Jesus' body (Matthew 27:62-66; 28:11-15).")
bullet("God cannot lie (Numbers 23:19; Hebrews 6:13-18), and He hates “a false "
       "tongue” (Proverbs 6:16-19). His law to the Israelites required "
       "compensation for injuries resulting from deception or malicious lying "
       "(Leviticus 6:2-7; 19:11-12).")
bullet("A person presenting false testimony was to receive the punishment that he "
       "desired to inflict upon another by means of his lies "
       "(Deuteronomy 19:15-21).")
bullet("God's view of malicious lying, as reflected in the law, has not changed. "
       "Those desiring to gain His approval cannot engage in the practice of lying "
       "(Proverbs 20:19; Colossians 3:9-10; 1 Timothy 3:11). They cannot be living "
       "a lie, claiming to love God while hating their brother (1 John 4:20-21).")
bullet("For playing false to the Holy Spirit by lying, Ananias and his wife lost "
       "their lives (Acts 5:1-11).")
bullet("However, persons who are momentarily overpowered in telling a lie do not "
       "automatically become guilty of an unforgivable sin, like Peter in denying "
       "Jesus three times (Matthew 26:69-75).")
bullet("While malicious lying is definitely condemned in the Bible, this does not "
       "mean that a person is under obligation to divulge truthful information to "
       "people who are not entitled to it. Jesus said we should not give what is "
       "holy to the dogs (Matthew 7:6).")
bullet("That is why Jesus, on certain occasions, refrained from giving full "
       "information or direct answers to certain questions, when doing so could "
       "have brought unnecessary harm (Matthew 15:1-6; 21:23-27; John 7:3-10).")
bullet("Evidently the course of Abraham, Isaac, Rahab, and Elisha in misdirecting "
       "or in withholding full facts from non-worshippers of God must be viewed "
       "in the same light (Genesis 12:10-19; Genesis 20; 26:1-10; Joshua 2:1-6; "
       "2 Kings 6:11-23).")
bullet("God allows “an operation of error” to go to persons who prefer "
       "falsehood, that they may get to believing the lie rather than the good "
       "news about Jesus Christ (2 Thessalonians 2:9-12).")

lead("This principle is illustrated by what happened centuries earlier in the "
     "case of the Israelite king Ahab. Lying prophets assured Ahab of success in "
     "war against Ramoth-gilead, while God's prophet Micaiah foretold disaster. As "
     "revealed in a vision to Micaiah, God allowed a “deceptive spirit” in the "
     "mouth of Ahab's prophets. That is to say, this evil spirit exercised his "
     "power upon them so that they spoke, not truth, but what they themselves "
     "wanted to say and what Ahab wanted to hear from them. Through this "
     "forewarning, Ahab preferred to be fooled by their lies and paid for it with "
     "his life (1 Kings 22:1-38; 2 Chronicles 18).")

# ================= SAVE =================
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Lie - draft.docx")
try:
    doc.save(out)
except PermissionError:
    raise SystemExit("Cannot write '%s'. It is open in Word. Close it and run again." % out)
print("Saved:", out)
print("Paragraphs:", len(doc.paragraphs))
