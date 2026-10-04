# -*- coding: utf-8 -*-
"""Build the Love chapter (typed sheets 1-6) as a Word doc matching the style of
'Introduction - final.docx' and the other chapters: Garamond, A5, 13pt justified.
References in plain parentheses, hyphens in verse ranges, no em dashes. The sheets
were typed with English points and the scripture in Twi; the Twi is REPLACED with
the English NIV (kept in the VERSES table below), keeping Dad's thirteen numbered
qualities of love and his structure. Dad's handwritten notes are worked in where
he placed them."""
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Garamond"
BODY_SZ = Pt(13)

doc = Document()

# ---- base style ----
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = BODY_SZ
rpr = normal.element.get_or_add_rPr()
rf = rpr.get_or_add_rFonts()
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

# ---- page = A5, margins from the first book ----
sec = doc.sections[0]
sec.page_width = Mm(148)
sec.page_height = Mm(210)
sec.top_margin = Mm(14)
sec.bottom_margin = Mm(15)
sec.left_margin = Mm(15)
sec.right_margin = Mm(13)

# ---- footer page number ----
def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p)

# ---- helpers ----
def _set_run(r, bold=False, italic=False, size=None):
    r.font.name = FONT
    rpr = r._r.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    r.font.bold = bold
    r.font.italic = italic
    if size: r.font.size = size

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text.upper())
    _set_run(r, bold=True, size=Pt(22))
    return p

def subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text); _set_run(r, italic=True, size=Pt(14))
    return p

def section_heading(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(15))
    return p

def textline(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(12))
    return p

def qheading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(13))
    return p

def lead(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); _set_run(r)
    return p

def bullet(text, keep=False, mark="•"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(1 if keep else 3)
    p.keep_with_next = keep
    r = p.add_run(mark + "  "); _set_run(r)
    r = p.add_run(text); _set_run(r)
    return p

def sub(text):
    """A short sub-list item (no reference), one level in."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Mm(12)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("–  "); _set_run(r)   # en dash marker, as in the Sabbath sub-lists
    r = p.add_run(text); _set_run(r)
    return p

def numbered(n, text, keep=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(8)
    p.paragraph_format.first_line_indent = Mm(-8)
    p.paragraph_format.space_after = Pt(1 if keep else 3)
    p.keep_with_next = keep
    r = p.add_run(("%s\t" if str(n).startswith("(") else "%s.\t") % n); _set_run(r)
    r = p.add_run(text); _set_run(r)
    return p

def scripture(ref, label=None):
    """The NIV text of `ref`, italic and indented, with the reference after it."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(10)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(7)
    USED.append(ref)
    text = VERSES.get(ref) or "[NIV TEXT PENDING]"
    r = p.add_run("“" + text + "”"); _set_run(r, italic=True, size=Pt(12))
    r = p.add_run(" (%s)" % (label or ref)); _set_run(r, size=Pt(12))
    return p

# ---- NIV text, keyed by reference (replaces Dad's Twi quotations) ----
VERSES = {
}
USED = []

# ================= CONTENT =================

title("Love")
textline("Text: 1 Corinthians 13:1-3")

bullet("The fruit of the Spirit is love (Galatians 5:22).")
lead("(John 13:34-35; 1 Corinthians 13:4).")
scripture("John 13:34-35")
scripture("1 Corinthians 13:4")
bullet("The one who loves has obeyed the law (Romans 13:8).")
bullet("The Spirit burns with love.")
bullet("Sanctify your souls with love (1 Peter 1:22-25).")

# ---- sheet 1 ----
qheading("1) A person must love God")
lead("(Deuteronomy 6:5).")
scripture("Deuteronomy 6:5")
bullet("His Son (Ephesians 6:24).", keep=True)
scripture("Ephesians 6:24")
bullet("And then the whole congregation (1 Peter 2:17; 1 John 2:10).", keep=True)
scripture("1 Peter 2:17")
scripture("1 John 2:10")
bullet("Love is a protector (1 Thessalonians 5:7-8).")
bullet("Let brotherly love continue (Hebrews 13:1).")

qheading("2) Love is long-suffering")
bullet("It puts up with unfavourable conditions and wrong actions of others, in "
       "order to help work out their salvation in Christ (Ephesians 4:32).",
       keep=True)
scripture("Ephesians 4:32")

qheading("3) Love is kind")
bullet("No matter what the provocation may be, harshness on the part of a "
       "Christian works no good (Romans 2:4; Titus 3:3-5).", keep=True)
scripture("Romans 2:4")
scripture("Titus 3:3-5")

qheading("4) Love is not jealous")
# sheet 2
bullet("It is not envious of good things coming to others.")
bullet("It rejoices in seeing a fellow man receive a position.")
bullet("It doesn't begrudge even one's enemies receiving good things "
       "(Matthew 5:45).", keep=True)
scripture("Matthew 5:45")
bullet("God's servants who have love are content with their lot "
       "(1 Timothy 6:6-8).", keep=True)
scripture("1 Timothy 6:6-8")
bullet("It was Satan who enviously desired to be worshipped (Luke 4:5-8).",
       keep=True)
scripture("Luke 4:5-8")

qheading("5) Love does not brag")
bullet("It doesn't seek the applause and admiration of creatures "
       "(Psalm 75:4-7).", keep=True)
scripture("Psalm 75:4-7")
bullet("He will not boast (Proverbs 27:1).", keep=True)
scripture("Proverbs 27:1")
bullet("He would know that all that he does is by God's grace (Psalm 34:2; "
       "Matthew 23:5-7).", keep=True)
scripture("Psalm 34:2")
bullet("A liar loves God but hates his brother (1 John 4:20-21).")

# ---- sheet 3 ----
qheading("6) Love doesn't get puffed up")
lead("(1 Corinthians 1:31; Jeremiah 9:24).")
scripture("1 Corinthians 1:31")
scripture("Jeremiah 9:24")
bullet("A person having love will not push another person down to make himself "
       "appear greater; rather, he will exalt God and will sincerely encourage and "
       "build up other persons (Colossians 1:3-5).", keep=True)
scripture("Colossians 1:3-5")

qheading("7) Love doesn't behave indecently")
bullet("It is not ill-mannered.")
bullet("It doesn't engage in indecent behaviour such as sexual abuse or shocking "
       "conduct.")
bullet("A person who has love will avoid doing things that disturb his fellow "
       "Christians (1 Corinthians 14:40).", keep=True)
scripture("1 Corinthians 14:40")
bullet("Love will prompt one to work honourably before non-Christians "
       "(Romans 13:13; 1 Thessalonians 4:12).", keep=True)
scripture("Romans 13:13")
scripture("1 Thessalonians 4:12")

qheading("8) Love doesn't look for its own interest")
bullet("Love applies the principle of 1 Corinthians 10:24.", keep=True)
scripture("1 Corinthians 10:24")
# sheet 4
bullet("Things are not always done in his way, with love (1 Corinthians 9:22-23).",
       keep=True)
scripture("1 Corinthians 9:22-23")
bullet("Love is more concerned about the spiritual welfare of others "
       "(Romans 14:13, 15).", keep=True)
scripture("Romans 14:13")
scripture("Romans 14:15")

qheading("9) Love doesn't become provoked")
bullet("It doesn't look for an excuse for provocation.")
bullet("It is not moved to outbursts of anger. It is the work of the flesh "
       "(Galatians 5:19-20).", keep=True)
scripture("Galatians 5:19-20")
bullet("Love is not easily offended by what others say or do.")
bullet("He is not afraid that his personal dignity may be injured.")

qheading("10) Love doesn't keep account of the injury")
bullet("It will not impute evil motives to another, but will give people the "
       "benefit of the doubt (Romans 14:1, 5).", keep=True)
scripture("Romans 14:1")
scripture("Romans 14:5")

qheading("11) Love doesn't rejoice over unrighteousness, but rejoices with the "
         "truth")
# sheet 5
bullet("It rejoices with the truth, even though it upsets previous beliefs or "
       "statements made.")
bullet("It sticks with God's word of truth.")
bullet("It sides with right, finding no pleasure in wrong, in lies, or in any form "
       "of injustice, no matter who the victim is, even if he is an enemy.")
bullet("It sees wrong to be wrong and fears not to speak against it "
       "(Galatians 2:11-14).", keep=True)
scripture("Galatians 2:11-14")
bullet("Also, it prefers to suffer wrong rather than to commit another wrong in "
       "order to straighten out the matter.")

qheading("12) Love bears all things")
bullet("It is willing to endure, to suffer for righteousness' sake.")
bullet("With love, one is unwilling to disclose to others those who have wronged "
       "him.")
bullet("If the offence is not serious, he will overlook it.")
bullet("If the offence is big, he will follow Matthew 18:15-17.", keep=True)
scripture("Matthew 18:15-17")
# sheet 6
scripture("Proverbs 10:12")
scripture("1 Peter 4:7-8")

qheading("13) Love believes all things")
bullet("Love has faith in the things God has said, even if outward appearances are "
       "against it and the unbelieving world scoffs (Joshua 23:14).", keep=True)
scripture("Joshua 23:14")
bullet("Love trusts in God's directions for the Christian congregation and His "
       "appointed servants (1 Timothy 5:17; Hebrews 13:17).", keep=True)
scripture("1 Timothy 5:17")
scripture("Hebrews 13:17")

# ================= SAVE =================
pending = sorted(set(r for r in USED if not VERSES.get(r)))
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Love - draft.docx")
try:
    doc.save(out)
except PermissionError:
    raise SystemExit("Cannot write '%s'. It is open in Word. Close it and run again." % out)
print("Saved:", out)
print("Paragraphs:", len(doc.paragraphs))
print("NIV text still pending for %d passages:" % len(pending), ", ".join(pending))
