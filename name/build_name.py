# -*- coding: utf-8 -*-
"""Build the Name chapter (handwritten sheets 1-5) as a Word doc matching the
style of 'Introduction - final.docx' and the other chapters: Garamond, A5, 13pt
justified. References in plain parentheses, hyphens in verse ranges, no em dashes.
The chapter runs under four headings, as Dad titled them: the meaning of the term,
its various uses, name as fame or reputation, and how one's name is damaged."""
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Garamond"
BODY_SZ = Pt(13)

doc = Document()

# ---- base style ----
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = BODY_SZ
rpr = normal.element.get_or_add_rPr()
rf = rpr.get_or_add_rFonts()
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

# ---- page = A5, margins from the first book ----
sec = doc.sections[0]
sec.page_width = Mm(148)
sec.page_height = Mm(210)
sec.top_margin = Mm(14)
sec.bottom_margin = Mm(15)
sec.left_margin = Mm(15)
sec.right_margin = Mm(13)

# ---- footer page number ----
def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p)

# ---- helpers ----
def _set_run(r, bold=False, italic=False, size=None):
    r.font.name = FONT
    rpr = r._r.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    r.font.bold = bold
    r.font.italic = italic
    if size: r.font.size = size

def _segments(p, segs):
    """Plain string, or a list of (text, italic) pairs for inline NIV quotes."""
    if isinstance(segs, str):
        segs = [(segs, False)]
    for txt, ital in segs:
        r = p.add_run(txt); _set_run(r, italic=ital)

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(text.upper())
    _set_run(r, bold=True, size=Pt(22))
    return p

def qheading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(13))
    return p

def lead(segs):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _segments(p, segs)
    return p

def bullet(segs):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("•  "); _set_run(r)
    _segments(p, segs)
    return p

# ================= CONTENT =================

title("Name")

# ---- sheet 1 ----
qheading("What is the meaning of the term “name”?")
lead("Name refers to a word or phrase that constitutes a distinctive designation "
     "of a person, place, animal, plant, or other object. “Name” can mean a "
     "person's reputation, or the person himself.")

bullet("Every family in heaven and on earth owes its name to God "
       "(Ephesians 3:14-15).")
bullet("An interesting example of how something completely new was named involves "
       "the miraculously provided manna. When the Israelites first saw it, they "
       "exclaimed, “What is it?” (Exodus 16:15). It was for this that they called "
       "it “manna”, probably meaning “What is it?” (Exodus 16:31).")

qheading("What are the various uses of the word “name”?")
bullet([("A particular name might be “called upon” a person, city, or building. "
         "Jacob, when adopting Joseph's two sons, Ephraim and Manasseh, as his own, "
         "stated: ", False),
        ("“May they be called by my name and the names of my fathers Abraham and "
         "Isaac”", True),
        (" (Genesis 48:16).", False)])
bullet("God's name being called on the Israelites indicated that they were His "
       "people (Deuteronomy 28:10; 2 Chronicles 7:14).")
# sheet 2
bullet("God also placed His name in Jerusalem and the temple, thereby accepting "
       "them as the rightful centre of His worship (2 Kings 21:4, 7).")
bullet("A person dying without leaving behind male offspring had his name “taken "
       "away”, as it were (Numbers 27:4). To forestall this situation, "
       "brother-in-law marriage was given by the law (Deuteronomy 25:5-6).")
bullet("On the other hand, the destruction of a nation, people, or family meant "
       "the wiping out of their name (Deuteronomy 7:24; 9:14; 1 Samuel 24:21).")
bullet("To speak or to act “in the name of” another denoted doing so as a "
       "representative of that one (Deuteronomy 10:8; Exodus 5:23).")
bullet("Similarly, to receive a person in the name of someone would indicate "
       "recognition of that one. Therefore, to “receive a prophet in the name of a "
       "prophet” would signify receiving a prophet because of his being such "
       "(Matthew 10:41).")
bullet("And to baptize in the name of Jesus would mean in recognition of Jesus.")
bullet("And to baptize in the name of the Father, the Son and the Holy Spirit would "
       "mean in recognition of the Father, the Son and the Holy Spirit "
       "(Matthew 28:19).")

# ---- sheet 3 ----
qheading("Can one's name mean his fame or reputation?")
bullet("In scriptural usage, “name” often denotes fame or reputation "
       "(1 Chronicles 14:17).")
bullet("Bringing a bad name upon someone meant making a false accusation against "
       "that person, marring his reputation (Deuteronomy 22:19).")
bullet("To have one's name “cast out as wicked” would mean a loss of good "
       "reputation (Luke 6:22).")
bullet("It was to make a celebrated “name” for themselves, in defiance of God, "
       "that men began building a tower and a city after Noah's Flood "
       "(Genesis 11:3-4).")
bullet("On the other hand, God promised to make Abraham's name great if he would "
       "leave his country and relatives to go to another land (Genesis 12:1-2).")
bullet([("At birth a person has no reputation, and therefore his name is little "
         "more than a label. That is why Ecclesiastes 7:1 says: ", False),
        ("“A good name is better than fine perfume, and the day of death better "
         "than the day of birth.”", True)])
bullet("Not at birth, but during the full course of a person's life, does his "
       "“name” take on real meaning, in the sense of identifying him either as a "
       "person practising righteousness or as one practising wickedness "
       "(Proverbs 22:1).")
bullet("By Jesus' faithfulness until death, His name became the one name “given "
       "among men by which we must get saved”, and He “inherited a name more "
       "excellent” than that of the angels (Acts 4:12; Hebrews 1:3-4).")
# sheet 4
bullet("But Solomon, for whom the hope was expressed that his name might become "
       "“more splendid” than David's, went into death with the name of a "
       "backslider as to true worship (1 Kings 1:47; 11:6, 9-11).")
bullet("“The very name of the wicked ones will rot”, or become an odious stench "
       "(Proverbs 10:7). For this reason, a good name “is to be chosen rather than "
       "abundant riches” (Proverbs 22:1).")

qheading("How one's name is damaged")
lead("In the Bible, a person's name is damaged or ruined through sin, a bad "
     "reputation, wickedness, or having one's name “blotted out” from God's book "
     "of records or life.")
bullet("Wicked behaviour and foolish choices are a sure way to destroy one's name. "
       "Proverbs 10:7 states that “the name of the wicked will rot”, meaning that "
       "a bad life leaves behind a foul legacy.")
bullet("Scandal and poor actions like pride, laziness, dishonesty, or harming "
       "others destroy a “good name” (Proverbs 22:1).")
# sheet 5
bullet("Past sins: even figures like the apostle Paul faced a damaged reputation "
       "initially, because his past name was synonymous with persecuting "
       "Christians.")

# ================= SAVE =================
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Name - draft.docx")
try:
    doc.save(out)
except PermissionError:
    raise SystemExit("Cannot write '%s'. It is open in Word. Close it and run again." % out)
print("Saved:", out)
print("Paragraphs:", len(doc.paragraphs))
