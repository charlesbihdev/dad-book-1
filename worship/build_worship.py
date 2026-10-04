# -*- coding: utf-8 -*-
"""Build the Worship chapter (handwritten sheets 1-2) as a Word doc matching the
style of 'Introduction - final.docx' and the other chapters: Garamond, A5, 13pt
justified. References in plain parentheses, hyphens in verse ranges, no em dashes.
Two headings, as Dad titled them: what worship is, and what unacceptable worship
is (four forms, numbered i to iv as Dad numbered them)."""
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Garamond"
BODY_SZ = Pt(13)

doc = Document()

# ---- base style ----
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = BODY_SZ
rpr = normal.element.get_or_add_rPr()
rf = rpr.get_or_add_rFonts()
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

# ---- page = A5, margins from the first book ----
sec = doc.sections[0]
sec.page_width = Mm(148)
sec.page_height = Mm(210)
sec.top_margin = Mm(14)
sec.bottom_margin = Mm(15)
sec.left_margin = Mm(15)
sec.right_margin = Mm(13)

# ---- footer page number ----
def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p)

# ---- helpers ----
def _set_run(r, bold=False, italic=False, size=None):
    r.font.name = FONT
    rpr = r._r.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    r.font.bold = bold
    r.font.italic = italic
    if size: r.font.size = size

def _segments(p, segs):
    """Plain string, or a list of (text, italic) pairs for inline NIV quotes."""
    if isinstance(segs, str):
        segs = [(segs, False)]
    for txt, ital in segs:
        r = p.add_run(txt); _set_run(r, italic=ital)

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(text.upper())
    _set_run(r, bold=True, size=Pt(22))
    return p

def qheading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.keep_with_next = True
    r = p.add_run(text); _set_run(r, bold=True, size=Pt(13))
    return p

def lead(segs):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _segments(p, segs)
    return p

def bullet(segs):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("•  "); _set_run(r)
    _segments(p, segs)
    return p

def form_point(label, text):
    """A numbered form of unacceptable worship: bold label, then Dad's text."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(8)
    p.paragraph_format.first_line_indent = Mm(-8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label); _set_run(r, bold=True)
    r = p.add_run(text); _set_run(r)
    return p

# ================= CONTENT =================

title("Worship")

# ---- sheet 1 ----
qheading("What is worship?")
lead("In the Bible, worship is an active, everyday lifestyle of honouring, obeying, "
     "and serving God with your whole heart and life. It is not just mentioning "
     "and singing to praise God's name.")

bullet("True worship of the Creator embraces every aspect of an individual's life "
       "(1 Corinthians 10:31).")
bullet("When God created Adam, He did not prescribe a particular ceremony by which "
       "man might approach Him in worship. Nevertheless, Adam was able to serve his "
       "Creator by faithfully doing the will of his heavenly Father.")
bullet("Later, to the nation of Israel, God did outline a certain way of approach "
       "in worship, including sacrifice, a priesthood, and a material sanctuary. "
       "This, however, had only “a shadow of the good things to come, but not the "
       "very substance of the things” (Hebrews 10:1).")
bullet("The primary emphasis in worship has always been on exercising faith and "
       "doing the will of God, and not on ceremony or ritual (Matthew 7:21; "
       "James 2:17-26).")
bullet("God, through the prophet Micah, provided the way for true worship "
       "(Micah 6:6-8).")

# ---- sheet 2 ----
qheading("What is unacceptable worship?")
lead("Unacceptable worship in the Bible is any religious service or praise that "
     "God rejects because it lacks sincere devotion, disobeys His commands, or is "
     "paired with unrepentant sin. Unacceptable worship comes in many forms. These "
     "include:")

form_point("(i)\tDisobedient worship: ",
           "This is the form in which a worshipper or worshippers do what they "
           "think is best instead of what God commanded, such as Cain's rejected "
           "offering in Genesis 4, or King Saul offering sacrifices in wilful "
           "disobedience (1 Samuel 15:22).")
form_point("(ii)\tHypocritical worship: ",
           "It means honouring God with outward words and songs while the heart is "
           "far from Him, which Jesus condemned, quoting Isaiah (Matthew 15:7-9).")
form_point("(iii)\tWorship mixed with iniquity: ",
           "Holding on to deliberate sin, injustice, or the oppression of the poor "
           "while attending solemn assemblies, which God calls an abomination "
           "(Isaiah 1:13; Amos 5:21-24).")
form_point("(iv)\tMan-made traditions: ",
           "This is where worshippers replace God's actual commandments with human "
           "customs or societal traditions (Mark 7:7-9).")

# ================= SAVE =================
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Worship - draft.docx")
try:
    doc.save(out)
except PermissionError:
    raise SystemExit("Cannot write '%s'. It is open in Word. Close it and run again." % out)
print("Saved:", out)
print("Paragraphs:", len(doc.paragraphs))
